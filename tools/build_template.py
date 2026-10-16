"""H-IFIT 원자료 입력 양식(.xlsx) 생성 스크립트 (v2).

사용: python3 build_template.py <출력경로.xlsx> [--sample]
  --sample : 검산용 가상 자료(경계 사례 포함)를 채운다. 배포본에는 쓰지 않는다.

배포본은 이 스크립트가 쓴 파일을 그대로 쓴다(LibreOffice로 다시 저장하지 않음).
수식 결과는 Excel이 파일을 열 때 계산한다(fullCalcOnLoad).
"""
import sys
import random
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
SAMPLE = "--sample" in sys.argv

# ------------------------------------------------------------------ 스타일
FONT = "맑은 고딕"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SEC = Font(name=FONT, size=11, bold=True)
F_H2 = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_VAR = Font(name=FONT, size=9, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_EX = Font(name=FONT, size=10, italic=True, color="7F7F7F")
F_NOTE = Font(name=FONT, size=9, color="595959")
F_RED = Font(name=FONT, size=10, bold=True, color="C00000")

FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_CALC = PatternFill("solid", fgColor="E7E6E6")
FILL_EX = PatternFill("solid", fgColor="F2F2F2")
FILL_HDR = PatternFill("solid", fgColor="1F4E78")
FILL_PRE = PatternFill("solid", fgColor="DDEBF7")
FILL_POST = PatternFill("solid", fgColor="E2EFDA")
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_WARN = PatternFill("solid", fgColor="F8CBAD")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_DIFF = PatternFill("solid", fgColor="BDD7EE")   # 포스터와 다름(중립색)

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MID = Alignment(horizontal="center", vertical="center")
LEFT_NOWRAP = Alignment(horizontal="left", vertical="center", wrap_text=False)

FIRST, LAST = 4, 33          # 자료 입력 행(30명분)
EX_ROW = 3                   # 예시 행

S_GUIDE, S_INFO, S_SUBJ, S_FMS, S_FMS2, S_ATT, S_PROG, S_FLOW, S_CALC, S_LOG = (
    "안내", "연구정보", "대상자_VO2max", "FMS", "FMS_평가자2",
    "출석_훈련기록", "주차별_프로그램", "대상자_흐름", "자동검산", "수정기록",
)


def q(s):
    return f"'{s}'"


wb = Workbook()
ws_guide = wb.active
ws_guide.title = S_GUIDE
ws_info, ws_subj, ws_fms, ws_fms2, ws_att, ws_prog, ws_flow, ws_calc, ws_log = (
    wb.create_sheet(s) for s in (S_INFO, S_SUBJ, S_FMS, S_FMS2, S_ATT, S_PROG, S_FLOW, S_CALC, S_LOG))


def add_dv(ws, rng, **kw):
    dv = DataValidation(**kw)
    dv.add(rng)
    ws.add_data_validation(dv)


def warn_cf(ws, col, first=FIRST, last=LAST):
    ws.conditional_formatting.add(f"{col}{first}:{col}{last}",
                                  FormulaRule(formula=[f"LEN({col}{first})>0"], fill=FILL_WARN))


def hdr_cell(ws, row, col, text, font=F_H2, fill=FILL_HDR):
    c = ws.cell(row=row, column=col, value=text)
    c.font, c.fill, c.alignment, c.border = font, fill, CENTER, BOX
    return c


def layout_data_sheet(ws, cols, formulas, label_height, freeze_col, examples=None):
    """1행 변수명, 2행 설명, 3행 예시, 4–33행 입력. cols: (name,label,kind,width,fmt)."""
    for i, (name, label, kind, width, fmt) in enumerate(cols, start=1):
        hdr_cell(ws, 1, i, name, font=F_VAR).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)
        c2 = ws.cell(row=2, column=i, value=label)
        c2.font, c2.alignment, c2.border = F_BOLD, CENTER, BOX
        c2.fill = {"pre": FILL_PRE, "post": FILL_POST, "calc": FILL_CALC}.get(kind, FILL_SEC)
        ws.column_dimensions[L(i)].width = width
        for row in [EX_ROW] + list(range(FIRST, LAST + 1)):
            c = ws.cell(row=row, column=i)
            c.border, c.number_format, c.alignment = BOX, fmt, MID
            if kind == "calc":
                c.value = formulas(row)[name]
                c.fill = FILL_CALC
                c.font = F_EX if row == EX_ROW else F_BASE
            elif row == EX_ROW:
                c.value = (examples or {}).get(name)
                c.fill, c.font = FILL_EX, F_EX
            else:
                c.fill, c.font = FILL_IN, F_INPUT
    ws.row_dimensions[1].height = 16
    ws.row_dimensions[2].height = label_height
    ws.freeze_panes = ws.cell(row=FIRST, column=freeze_col)


def col_map(cols):
    return {name: L(i + 1) for i, (name, *_r) in enumerate(cols)}


# ================================================================== 대상자_VO2max
EXCL_LIST = ["중재 중단", "사후 미측정", "부상", "최대치 판정 미충족", "기타"]
COLS_S = [
    ("ID", "대상자 ID\n(예: E01, C01)", "in", 9, "@"),
    ("group", "집단\n1=실험\n2=통제", "in", 7, "0"),
    ("group_lbl", "집단명\n(자동)", "calc", 8, "@"),
    ("include", "분석 포함\n1=포함 0=제외\n(FMS는\n기본값)", "in", 9, "0"),
    ("excl_reason", "제외 사유\n(목록 선택)", "in", 14, "@"),
    ("excl_detail", "제외 사유\n상세", "in", 16, "@"),
    ("team", "소속 팀·\n대학 코드", "in", 9, "@"),
    ("sex", "성별\n1=남\n2=여", "in", 6, "0"),
    ("age", "연령\n(만, 세)", "in", 7, "0"),
    ("height", "신장\n(cm)", "in", 7, "0.0"),
    ("wclass", "체급\n(예: -68kg)", "in", 9, "@"),
    ("career", "선수 경력\n(년)", "in", 7, "0.0"),
    ("level", "등록 수준·단·\n주요 입상", "in", 13, "@"),
    ("wt_pre", "사전 체중\n(kg)", "pre", 8, "0.0"),
    ("vo2_pre", "사전 VO2max\n(mL/kg/min)", "pre", 10, "0.00"),
    ("vo2abs_pre", "사전 VO2max\n절대값\n(mL/min)", "pre", 9, "0"),
    ("hrmax_pre", "사전 검사\n최고심박\n(bpm)", "pre", 8, "0"),
    ("rer_pre", "사전 검사\n최고 RER", "pre", 8, "0.00"),
    ("rpe_pre", "사전 검사\n종료 RPE\n(6–20)", "pre", 8, "0"),
    ("plateau_pre", "사전 VO2\nplateau\n1=예 0=아니오", "pre", 9, "0"),
    ("crit_pre", "사전 최대치\n판정 충족\n1=예 0=아니오", "pre", 9, "0"),
    ("date_pre", "사전 VO2max\n검사일\n(YYYY-MM-DD)", "pre", 11, "yyyy-mm-dd"),
    ("wt_post", "사후 체중\n(kg)", "post", 8, "0.0"),
    ("vo2_post", "사후 VO2max\n(mL/kg/min)", "post", 10, "0.00"),
    ("vo2abs_post", "사후 VO2max\n절대값\n(mL/min)", "post", 9, "0"),
    ("hrmax_post", "사후 검사\n최고심박\n(bpm)", "post", 8, "0"),
    ("rer_post", "사후 검사\n최고 RER", "post", 8, "0.00"),
    ("rpe_post", "사후 검사\n종료 RPE\n(6–20)", "post", 8, "0"),
    ("plateau_post", "사후 VO2\nplateau\n1=예 0=아니오", "post", 9, "0"),
    ("crit_post", "사후 최대치\n판정 충족\n1=예 0=아니오", "post", 9, "0"),
    ("date_post", "사후 VO2max\n검사일\n(YYYY-MM-DD)", "post", 11, "yyyy-mm-dd"),
    ("d_vo2", "VO2max 변화\n사후−사전\n(자동)", "calc", 9, "0.00"),
    ("pct_vo2", "VO2max\n변화율\n(자동)", "calc", 8, "0.0%"),
    ("d_wt", "체중 변화\n(kg, 자동)", "calc", 8, "0.0"),
    ("critn_pre", "사전 충족\n기준 수\n(0–4, 자동)", "calc", 8, "0"),
    ("critn_post", "사후 충족\n기준 수\n(0–4, 자동)", "calc", 8, "0"),
    ("chk_vo2", "자동 점검 (경고가 있으면 주황색)", "calc", 60, "@"),
]
CS = col_map(COLS_S)
EX_S = {"ID": "예시", "group": 1, "include": 0, "excl_reason": "기타", "excl_detail": "예시 행(분석 제외)",
        "team": "A대", "sex": 1, "age": 20, "height": 176.4, "wclass": "-68kg", "career": 8.5,
        "level": "대학부 등록, 3단",
        "wt_pre": 68.2, "vo2_pre": 48.30, "vo2abs_pre": 3294, "hrmax_pre": 196, "rer_pre": 1.15,
        "rpe_pre": 19, "plateau_pre": 1, "crit_pre": 1, "date_pre": dt.date(2025, 3, 3),
        "wt_post": 67.5, "vo2_post": 52.40, "vo2abs_post": 3537, "hrmax_post": 195, "rer_post": 1.13,
        "rpe_post": 19, "plateau_post": 0, "crit_post": 1, "date_post": dt.date(2025, 5, 30)}


def critn(rer, hr, rpe, plat, age):
    return (f'=IF(COUNT({rer},{hr},{rpe},{plat})=0,"",'
            f'IF(ISNUMBER({rer}),--({rer}>=1.1),0)'
            f'+IF(AND(ISNUMBER({hr}),ISNUMBER({age})),--({hr}>=0.95*(220-{age})),0)'
            f'+IF(ISNUMBER({rpe}),--({rpe}>=17),0)'
            f'+IF(ISNUMBER({plat}),--({plat}=1),0))')


def not_num(ref, msg):
    return f'IF(AND({ref}<>"",NOT(ISNUMBER({ref}))),"{msg} 숫자 아님 ","")'


def rng_chk(ref, lo, hi, msg):
    return f'IF(ISNUMBER({ref}),IF(OR({ref}<{lo},{ref}>{hi}),"{msg} ",""),"")'


def formulas_subj(r):
    c = {k: f"{v}{r}" for k, v in CS.items()}
    ids = f"${CS['ID']}${FIRST}:${CS['ID']}${LAST}"
    parts = [
        f'IF(COUNTIF({ids},{c["ID"]})>1,"ID 중복 ","")',
        f'IF(LEN({c["ID"]})<>LEN(TRIM({c["ID"]})),"ID 앞뒤 공백 ","")',
        f'IF(AND({c["group"]}<>1,{c["group"]}<>2),"집단코드 확인 ","")',
        f'IF(OR({c["include"]}="",AND({c["include"]}<>0,{c["include"]}<>1)),"분석포함 코드 확인(빈칸=제외됨) ","")',
        f'IF(AND(ISNUMBER({c["include"]}),{c["include"]}=0,{c["excl_reason"]}=""),"제외 사유 없음 ","")',
        not_num(c["vo2_pre"], "사전VO2"), not_num(c["vo2_post"], "사후VO2"),
        not_num(c["wt_pre"], "사전체중"), not_num(c["wt_post"], "사후체중"),
        rng_chk(c["vo2_pre"], 25, 85, "사전VO2 범위(25–85) 확인"),
        rng_chk(c["vo2_post"], 25, 85, "사후VO2 범위(25–85) 확인"),
        f'IF(AND(ISNUMBER({c["vo2_pre"]}),ISNUMBER({c["vo2_post"]})),'
        f'IF(ABS({c["vo2_post"]}-{c["vo2_pre"]})>12,"VO2 변화량 12 초과 확인 ",""),"")',
        f'IF(AND(ISNUMBER({c["vo2abs_pre"]}),ISNUMBER({c["wt_pre"]}),ISNUMBER({c["vo2_pre"]})),'
        f'IF({c["wt_pre"]}<=0,"사전 체중 확인 ",IF(ABS({c["vo2abs_pre"]}/{c["wt_pre"]}-{c["vo2_pre"]})>0.5,'
        f'"사전 절대값/체중≠상대값 ","")),"")',
        f'IF(AND(ISNUMBER({c["vo2abs_post"]}),ISNUMBER({c["wt_post"]}),ISNUMBER({c["vo2_post"]})),'
        f'IF({c["wt_post"]}<=0,"사후 체중 확인 ",IF(ABS({c["vo2abs_post"]}/{c["wt_post"]}-{c["vo2_post"]})>0.5,'
        f'"사후 절대값/체중≠상대값 ","")),"")',
        rng_chk(c["rer_pre"], 0.85, 1.5, "사전 RER 범위 확인"),
        rng_chk(c["rer_post"], 0.85, 1.5, "사후 RER 범위 확인"),
        rng_chk(c["hrmax_pre"], 150, 225, "사전 최고심박 범위 확인"),
        rng_chk(c["hrmax_post"], 150, 225, "사후 최고심박 범위 확인"),
        f'IF(AND(ISNUMBER({c["crit_pre"]}),ISNUMBER({c["critn_pre"]})),'
        f'IF(AND({c["crit_pre"]}=1,{c["critn_pre"]}<2),"사전 판정충족=1이나 기준 2개 미만 ",""),"")',
        f'IF(AND(ISNUMBER({c["crit_post"]}),ISNUMBER({c["critn_post"]})),'
        f'IF(AND({c["crit_post"]}=1,{c["critn_post"]}<2),"사후 판정충족=1이나 기준 2개 미만 ",""),"")',
        f'IF(AND({c["date_pre"]}<>"",NOT(ISNUMBER({c["date_pre"]}))),"사전 검사일 날짜형식 아님 ","")',
        f'IF(AND({c["date_post"]}<>"",NOT(ISNUMBER({c["date_post"]}))),"사후 검사일 날짜형식 아님 ","")',
        f'IF(AND(ISNUMBER({c["date_pre"]}),ISNUMBER({c["date_post"]})),'
        f'IF({c["date_post"]}<={c["date_pre"]},"검사일 역전 ",'
        f'IF(OR({c["date_post"]}-{c["date_pre"]}<77,{c["date_post"]}-{c["date_pre"]}>126),"사전–사후 간격(11–18주) 확인 ","")),"")',
        f'IF({c["include"]}=1,IF(OR(NOT(ISNUMBER({c["vo2_pre"]})),NOT(ISNUMBER({c["vo2_post"]}))),"분석포함인데 VO2 결측 ",""),"")',
    ]
    return {
        "group_lbl": f'=IF({c["group"]}="","",IF({c["group"]}=1,"실험군",IF({c["group"]}=2,"통제군","코드오류")))',
        "d_vo2": f'=IF(AND(ISNUMBER({c["vo2_pre"]}),ISNUMBER({c["vo2_post"]})),{c["vo2_post"]}-{c["vo2_pre"]},"")',
        "pct_vo2": (f'=IF(AND(ISNUMBER({c["vo2_pre"]}),ISNUMBER({c["vo2_post"]})),'
                    f'IF({c["vo2_pre"]}<>0,({c["vo2_post"]}-{c["vo2_pre"]})/{c["vo2_pre"]},""),"")'),
        "d_wt": f'=IF(AND(ISNUMBER({c["wt_pre"]}),ISNUMBER({c["wt_post"]})),{c["wt_post"]}-{c["wt_pre"]},"")',
        "critn_pre": critn(c["rer_pre"], c["hrmax_pre"], c["rpe_pre"], c["plateau_pre"], c["age"]),
        "critn_post": critn(c["rer_post"], c["hrmax_post"], c["rpe_post"], c["plateau_post"], c["age"]),
        "chk_vo2": f'=IF({c["ID"]}="","",TRIM({"&".join(parts)}))',
    }


layout_data_sheet(ws_subj, COLS_S, formulas_subj, 58, 2, EX_S)
for row in range(EX_ROW, LAST + 1):
    ws_subj[f"{CS['chk_vo2']}{row}"].alignment = LEFT_NOWRAP
for name in ("excl_reason", "excl_detail", "level"):
    for row in range(EX_ROW, LAST + 1):
        ws_subj[f"{CS[name]}{row}"].alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
ws_subj[f"A{EX_ROW}"].comment = Comment(
    "예시 행입니다. 분석포함=0이고 자료 범위(4–33행) 밖이라 지우지 않아도 자동검산에서 빠집니다. "
    "행을 삽입·삭제하지 말고, 탈락자도 행을 지우지 말고 분석포함=0과 사유를 적으세요.", "양식", width=380, height=200)
srng = lambda n: f"{CS[n]}{FIRST}:{CS[n]}{LAST}"
add_dv(ws_subj, srng("group"), type="list", formula1='"1,2"', allow_blank=True, showErrorMessage=True,
       errorTitle="집단 코드", error="1(실험군) 또는 2(통제군)만 입력하세요.")
add_dv(ws_subj, srng("include"), type="list", formula1='"1,0"', allow_blank=True, showErrorMessage=True,
       errorTitle="분석 포함", error="1(포함) 또는 0(제외)만 입력하세요.")
add_dv(ws_subj, srng("excl_reason"), type="list", formula1='"' + ",".join(EXCL_LIST) + '"', allow_blank=True,
       showErrorMessage=True, errorStyle="information", errorTitle="제외 사유",
       error="목록에서 고르고, 자세한 내용은 오른쪽 '상세' 칸에 적어 주세요.")
add_dv(ws_subj, srng("sex"), type="list", formula1='"1,2"', allow_blank=True, showErrorMessage=True,
       errorTitle="성별", error="1(남) 또는 2(여)만 입력하세요.")
for n in ("crit_pre", "crit_post", "plateau_pre", "plateau_post"):
    add_dv(ws_subj, srng(n), type="list", formula1='"1,0"', allow_blank=True, showErrorMessage=True,
           errorTitle="예/아니오", error="1(예) 또는 0(아니오)만 입력하세요. 판정할 수 없으면 빈칸.")
for n in ("vo2_pre", "vo2_post"):
    add_dv(ws_subj, srng(n), type="decimal", operator="between", formula1="15", formula2="95", allow_blank=True,
           showErrorMessage=True, errorStyle="warning", errorTitle="VO2max 단위",
           error="mL/kg/min 단위가 맞는지 확인하세요(일반 범위 15–95).")
for n in ("rpe_pre", "rpe_post"):
    add_dv(ws_subj, srng(n), type="decimal", operator="between", formula1="6", formula2="20", allow_blank=True,
           showErrorMessage=True, errorStyle="warning", errorTitle="RPE",
           error="Borg 6–20 기준입니다. 다른 척도면 연구정보 E5에 적어 주세요.")
for n in ("date_pre", "date_post"):
    add_dv(ws_subj, srng(n), type="date", operator="between", formula1="DATE(2020,1,1)",
           formula2="DATE(2027,12,31)", allow_blank=True, showErrorMessage=True, errorStyle="warning",
           errorTitle="날짜 형식", error="2025-03-03 형식(하이픈)으로 입력하세요. 2025.03.03은 날짜로 인식되지 않습니다.")
warn_cf(ws_subj, CS["chk_vo2"])

# ================================================================== FMS (평가자1) / FMS_평가자2 공통
ITEMS = [  # 코드, 이름, 양측, clearing 설명
    ("DS", "Deep Squat", False, None),
    ("HS", "Hurdle Step", True, None),
    ("ILL", "In-line Lunge", True, None),
    ("SM", "Shoulder Mobility", True, "어깨 충돌 검사(좌·우 중 한쪽이라도 통증이면 1)"),
    ("ASLR", "Active Straight Leg Raise", True, None),
    ("TSPU", "Trunk Stability Push-up", False, "척추 신전(prone press-up) 검사"),
    ("RS", "Rotary Stability", True, "척추 굴곡(posterior rocking) 검사"),
]


def move_names(t):
    out = []
    for code, _, bilat, _c in ITEMS:
        out += [f"{code}_L_{t}", f"{code}_R_{t}"] if bilat else [f"{code}_{t}"]
    return out


def clr_names(t):
    return [f"{code}_clr_{t}" for code, _, _b, clr in ITEMS if clr]


def raw_cols(t, with_extra=True):
    tk = "사전" if t == "pre" else "사후"
    cols = [(f"rater_{t}", f"{tk}\n평가자\n(코드)", t, 7, "@")]
    if with_extra:
        cols.append((f"date_fms_{t}", f"{tk} FMS\n검사일", t, 10, "yyyy-mm-dd"))
    for code, _, bilat, clr in ITEMS:
        if bilat:
            for side, sk in (("L", "좌"), ("R", "우")):
                cols.append((f"{code}_{side}_{t}", f"{tk}\n{code} {sk}\n(0–3)", t, 7, "0"))
        else:
            cols.append((f"{code}_{t}", f"{tk}\n{code}\n(0–3)", t, 7, "0"))
        if clr:
            extra = "\n(좌·우 중 하나라도)" if code == "SM" else ""
            cols.append((f"{code}_clr_{t}", f"{tk}\n{code}\nclearing\n통증 1/0{extra}", t, 9, "0"))
    if with_extra:
        cols.append((f"fms_spss_{t}", f"{tk} 총점\n기존 SPSS·\n포스터 값\n(있으면)", t, 9, "0.00"))
    return cols


def calc_cols(t, with_src=True):
    tk = "사전" if t == "pre" else "사후"
    cols = [(f"{code}s_{t}", f"{tk}\n{code}\n점수\n(자동)", "calc", 6, "0") for code, *_ in ITEMS]
    cols.append((f"fms_calc_{t}", f"{tk}\n하위항목 합\n(자동)", "calc", 8, "0"))
    cols.append((f"asym_{t}", f"{tk}\n좌우 비대칭\n항목 수\n(자동)", "calc", 8, "0"))
    if with_src:
        cols.append((f"clrn_{t}", f"{tk}\nclearing\n기록 수\n(자동)", "calc", 7, "0"))
    return cols


def lookup(col_name, r, id_col="A"):
    rng = f"{q(S_SUBJ)}!${CS[col_name]}${FIRST}:${CS[col_name]}${LAST}"
    ids = f"{q(S_SUBJ)}!${CS['ID']}${FIRST}:${CS['ID']}${LAST}"
    idx = f"INDEX({rng},MATCH(${id_col}{r},{ids},0))"
    return f'=IF(${id_col}{r}="","",IFERROR(IF({idx}="","",{idx}),"ID없음"))'


def fms_item_formulas(cmap, t, r):
    """하위항목 점수: 양측은 낮은 쪽, clearing 통증(1)이면 0, clearing 빈칸은 통증 없음."""
    f = {}
    items = []
    for code, _, bilat, clr in ITEMS:
        if bilat:
            a, b = f"{cmap[f'{code}_L_{t}']}{r}", f"{cmap[f'{code}_R_{t}']}{r}"
            core, cond = f"MIN({a},{b})", f"AND(ISNUMBER({a}),ISNUMBER({b}))"
        else:
            a = f"{cmap[f'{code}_{t}']}{r}"
            core, cond = a, f"ISNUMBER({a})"
        if clr:
            core = f"IF({cmap[f'{code}_clr_{t}']}{r}=1,0,{core})"
        f[f"{code}s_{t}"] = f'=IF({cond},{core},"")'
        items.append(f"{cmap[f'{code}s_{t}']}{r}")
    rng = f"{items[0]}:{items[-1]}"
    f[f"fms_calc_{t}"] = f'=IF(COUNT({rng})=7,SUM({rng}),"")'
    pairs = "+".join(f"--({cmap[f'{code}_L_{t}']}{r}<>{cmap[f'{code}_R_{t}']}{r})"
                     for code, _, bilat, _c in ITEMS if bilat)
    f[f"asym_{t}"] = f'=IF(ISNUMBER({cmap[f"fms_calc_{t}"]}{r}),{pairs},"")'
    if f"clrn_{t}" in cmap:
        clr = ",".join(f"{cmap[n]}{r}" for n in clr_names(t))
        f[f"clrn_{t}"] = f'=IF(ISNUMBER({cmap[f"fms_calc_{t}"]}{r}),COUNT({clr}),"")'
    return f


def partial_blank(cmap, t, r, label):
    mv = ",".join(f"{cmap[n]}{r}" for n in move_names(t))
    return f'IF(AND(COUNT({mv})>0,COUNT({mv})<12),"{label} 동작점수 일부 빈칸 ","")'


# ------------------------------------------------------------------ FMS 시트
COLS_F = ([("ID", "대상자 ID\n(직접 입력)", "in", 9, "@"),
           ("group", "집단\n(자동)", "calc", 6, "0"),
           ("include", "분석\n포함\n(자동)", "calc", 6, "0"),
           ("include_fms", "FMS만\n분석포함\n1/0 (비우면\n대상자 시트\n따름)", "in", 9, "0")]
          + raw_cols("pre") + raw_cols("post") + calc_cols("pre") + calc_cols("post")
          + [("fms_pre", "분석용\n사전 총점\n(자동)", "calc", 8, "0.00"),
             ("fms_post", "분석용\n사후 총점\n(자동)", "calc", 8, "0.00"),
             ("d_fms", "FMS 변화\n사후−사전\n(자동)", "calc", 8, "0.00"),
             ("src_pre", "사전 총점\n출처\n(자동)", "calc", 9, "@"),
             ("src_post", "사후 총점\n출처\n(자동)", "calc", 9, "@"),
             ("r2_pre", "평가자2\n사전 총점\n(자동)", "calc", 8, "0"),
             ("r2_post", "평가자2\n사후 총점\n(자동)", "calc", 8, "0"),
             ("avg_pre", "평가자 평균\n사전 총점\n(자동)", "calc", 9, "0.00"),
             ("avg_post", "평가자 평균\n사후 총점\n(자동)", "calc", 9, "0.00"),
             ("d_avg", "평가자 평균\n변화\n(자동)", "calc", 9, "0.00"),
             ("chk_fms", "자동 점검 (경고가 있으면 주황색)", "calc", 60, "@")])
CF = col_map(COLS_F)

COLS_F2 = ([("ID", "대상자 ID\n(직접 입력)", "in", 9, "@"),
            ("group", "집단\n(자동)", "calc", 6, "0")]
           + raw_cols("pre", with_extra=False) + raw_cols("post", with_extra=False)
           + calc_cols("pre", with_src=False) + calc_cols("post", with_src=False)
           + [("chk_fms2", "자동 점검 (경고가 있으면 주황색)", "calc", 50, "@")])
CF2 = col_map(COLS_F2)


def formulas_fms(r):
    c = {k: f"{v}{r}" for k, v in CF.items()}
    f = {"group": lookup("group", r),
         "include": f'=IF(ISNUMBER({c["include_fms"]}),{c["include_fms"]},{lookup("include", r)[1:]})'}
    if r == EX_ROW:  # 예시 행: 연결 대상이 없으므로 고정값
        f.update({"group": 1, "include": 0})
    for t in ("pre", "post"):
        f.update(fms_item_formulas(CF, t, r))
        calc, spss = c[f"fms_calc_{t}"], c[f"fms_spss_{t}"]
        f[f"fms_{t}"] = f'=IF(ISNUMBER({calc}),{calc},IF(ISNUMBER({spss}),{spss},""))'
        f[f"src_{t}"] = f'=IF(ISNUMBER({calc}),"하위항목",IF(ISNUMBER({spss}),"SPSS총점",""))'
        r2rng = f"{q(S_FMS2)}!${CF2[f'fms_calc_{t}']}${FIRST}:${CF2[f'fms_calc_{t}']}${LAST}"
        r2ids = f"{q(S_FMS2)}!${CF2['ID']}${FIRST}:${CF2['ID']}${LAST}"
        idx = f"INDEX({r2rng},MATCH($A{r},{r2ids},0))"
        f[f"r2_{t}"] = f'=IF($A{r}="","",IFERROR(IF(ISNUMBER({idx}),{idx},""),""))'
        f[f"avg_{t}"] = f'=IF(AND(ISNUMBER({calc}),ISNUMBER({c[f"r2_{t}"]})),({calc}+{c[f"r2_{t}"]})/2,"")'
    f["d_fms"] = f'=IF(AND(ISNUMBER({c["fms_pre"]}),ISNUMBER({c["fms_post"]})),{c["fms_post"]}-{c["fms_pre"]},"")'
    f["d_avg"] = f'=IF(AND(ISNUMBER({c["avg_pre"]}),ISNUMBER({c["avg_post"]})),{c["avg_post"]}-{c["avg_pre"]},"")'
    ids = f"${CF['ID']}${FIRST}:${CF['ID']}${LAST}"
    parts = [
        f'IF({c["group"]}="ID없음","대상자_VO2max 시트에 없는 ID ","")',
        f'IF(LEN({c["ID"]})<>LEN(TRIM({c["ID"]})),"ID 앞뒤 공백 ","")',
        f'IF(COUNTIF({ids},{c["ID"]})>1,"ID 중복 ","")',
        f'IF(AND({c["group"]}<>"ID없음",OR({c["group"]}="",AND({c["group"]}<>1,{c["group"]}<>2))),"집단 코드 확인(대상자 시트) ","")',
        f'IF(AND({c["group"]}<>"ID없음",OR({c["include"]}="",AND({c["include"]}<>0,{c["include"]}<>1))),"분석포함 코드 확인 ","")',
        partial_blank(CF, "pre", r, "사전"), partial_blank(CF, "post", r, "사후"),
        not_num(c["fms_spss_pre"], "사전 SPSS총점"), not_num(c["fms_spss_post"], "사후 SPSS총점"),
    ]
    for t, tk in (("pre", "사전"), ("post", "사후")):
        calc, spss, use = c[f"fms_calc_{t}"], c[f"fms_spss_{t}"], c[f"fms_{t}"]
        parts += [
            f'IF(AND(ISNUMBER({calc}),ISNUMBER({spss})),IF({calc}<>{spss},"{tk}: 계산총점≠SPSS총점 ",""),"")',
            f'IF(ISNUMBER({use}),IF({use}<>INT({use}),"{tk} 총점 정수 아님 ",""),"")',
            f'IF(ISNUMBER({use}),IF(OR({use}<0,{use}>21),"{tk} 총점 범위(0–21) 밖 ",""),"")',
        ]
    parts.append(f'IF({c["include"]}=1,IF(OR(NOT(ISNUMBER({c["fms_pre"]})),NOT(ISNUMBER({c["fms_post"]}))),'
                 f'"분석포함인데 FMS 결측 ",""),"")')
    f["chk_fms"] = "" if r == EX_ROW else f'=IF({c["ID"]}="","",TRIM({"&".join(parts)}))'
    return f


def formulas_fms2(r):
    c = {k: f"{v}{r}" for k, v in CF2.items()}
    f = {"group": 1 if r == EX_ROW else lookup("group", r)}
    for t in ("pre", "post"):
        f.update(fms_item_formulas(CF2, t, r))
    ids = f"${CF2['ID']}${FIRST}:${CF2['ID']}${LAST}"
    f1ids = f"{q(S_FMS)}!${CF['ID']}${FIRST}:${CF['ID']}${LAST}"
    parts = [
        f'IF({c["group"]}="ID없음","대상자_VO2max 시트에 없는 ID ","")',
        f'IF(LEN({c["ID"]})<>LEN(TRIM({c["ID"]})),"ID 앞뒤 공백 ","")',
        f'IF(COUNTIF({ids},{c["ID"]})>1,"ID 중복 ","")',
        f'IF(COUNTIF({f1ids},{c["ID"]})=0,"FMS(평가자1) 시트에 없는 ID ","")',
        partial_blank(CF2, "pre", r, "사전"), partial_blank(CF2, "post", r, "사후"),
    ]
    f["chk_fms2"] = "" if r == EX_ROW else f'=IF({c["ID"]}="","",TRIM({"&".join(parts)}))'
    return f


EX_PRE = {"DS": 2, "HS_L": 2, "HS_R": 2, "ILL_L": 2, "ILL_R": 3, "SM_L": 3, "SM_R": 2, "SM_clr": 0,
          "ASLR_L": 2, "ASLR_R": 2, "TSPU": 2, "TSPU_clr": 0, "RS_L": 2, "RS_R": 2, "RS_clr": 0}
EX_POST = {"DS": 3, "HS_L": 2, "HS_R": 3, "ILL_L": 3, "ILL_R": 3, "SM_L": 3, "SM_R": 2, "SM_clr": 0,
           "ASLR_L": 3, "ASLR_R": 2, "TSPU": 2, "TSPU_clr": 0, "RS_L": 2, "RS_R": 2, "RS_clr": 0}
EX_F = {"ID": "예시", "rater_pre": "R1", "rater_post": "R1", "date_fms_pre": dt.date(2025, 3, 4),
        "date_fms_post": dt.date(2025, 5, 31), "fms_spss_pre": 14, "fms_spss_post": 16}
EX_F.update({f"{k}_pre": v for k, v in EX_PRE.items()})
EX_F.update({f"{k}_post": v for k, v in EX_POST.items()})
EX_F2 = {k: v for k, v in EX_F.items() if k in CF2}
EX_F2.update({"rater_pre": "R2", "rater_post": "R2", "ILL_R_pre": 2, "DS_post": 2})

for ws, cols, fn, cmap, ex in ((ws_fms, COLS_F, formulas_fms, CF, EX_F),
                               (ws_fms2, COLS_F2, formulas_fms2, CF2, EX_F2)):
    layout_data_sheet(ws, cols, fn, 84, 5 if ws is ws_fms else 3, ex)
    chk = "chk_fms" if ws is ws_fms else "chk_fms2"
    for row in range(EX_ROW, LAST + 1):
        ws[f"{cmap[chk]}{row}"].alignment = LEFT_NOWRAP
    for t in ("pre", "post"):
        for n in move_names(t):
            add_dv(ws, f"{cmap[n]}{FIRST}:{cmap[n]}{LAST}", type="decimal", operator="between",
                   formula1="0", formula2="3", allow_blank=True, showErrorMessage=True, errorStyle="warning",
                   errorTitle="FMS 점수", error="FMS 원점수는 보통 0–3 정수입니다. 평가자 평균 등으로 소수라면 "
                   "그대로 입력하고 연구정보 F5에 적어 주세요.")
        for n in clr_names(t):
            add_dv(ws, f"{cmap[n]}{FIRST}:{cmap[n]}{LAST}", type="list", formula1='"1,0"', allow_blank=True,
                   showErrorMessage=True, errorTitle="Clearing test",
                   error="1(통증 있음) 또는 0(없음). 검사하지 않았거나 기록이 없으면 빈칸.")
        if t == "pre" and "include_fms" in cmap:
            add_dv(ws, f"{cmap['include_fms']}{FIRST}:{cmap['include_fms']}{LAST}", type="list", formula1='"1,0"',
                   allow_blank=True, showErrorMessage=True, errorTitle="FMS 분석포함",
                   error="1(포함) 또는 0(제외). 비우면 대상자 시트의 분석포함을 따릅니다.")
        if f"fms_spss_{t}" in cmap:
            add_dv(ws, f"{cmap[f'fms_spss_{t}']}{FIRST}:{cmap[f'fms_spss_{t}']}{LAST}", type="decimal",
                   operator="between", formula1="0", formula2="21", allow_blank=True, showErrorMessage=True,
                   errorTitle="FMS 총점", error="FMS 총점은 0–21 범위입니다.")
        if f"date_fms_{t}" in cmap:
            add_dv(ws, f"{cmap[f'date_fms_{t}']}{FIRST}:{cmap[f'date_fms_{t}']}{LAST}", type="date",
                   operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2027,12,31)", allow_blank=True,
                   showErrorMessage=True, errorStyle="warning", errorTitle="날짜 형식",
                   error="2025-03-03 형식(하이픈)으로 입력하세요.")
    warn_cf(ws, cmap[chk])
    ws["A1"].comment = Comment(
        "약어: DS=Deep Squat, HS=Hurdle Step, ILL=In-line Lunge, SM=Shoulder Mobility, "
        "ASLR=Active Straight Leg Raise, TSPU=Trunk Stability Push-up, RS=Rotary Stability. "
        "ID는 대상자_VO2max 시트와 같은 코드로 직접 입력합니다(행 순서가 달라도 ID로 연결).", "양식", width=380, height=200)
ws_fms[f"{CF['fms_spss_pre']}2"].comment = Comment(
    "하위항목 원점수가 없고 총점만 있을 때, 또는 기존 SPSS 파일의 총점과 대조할 때 입력. "
    "하위항목이 모두 있으면 하위항목 합이 우선 쓰이고, 다르면 경고가 뜹니다. "
    "자동검산 J에 SPSS총점으로 대체된 값의 수가 표시됩니다.", "양식", width=380, height=200)
ws_fms[f"{CF['SM_clr_pre']}2"].comment = Comment(
    "Clearing test: 1=통증 있음 → 해당 항목 0점, 0=통증 없음. 검사하지 않았거나 기록이 없으면 빈칸 "
    "(통증 없음으로 계산하고 자동검산 J에 개수 표시). "
    "SM=어깨 충돌 검사(좌·우 중 한쪽이라도 통증이면 1), TSPU=척추 신전(prone press-up), "
    "RS=척추 굴곡(posterior rocking).", "양식", width=380, height=200)
ws_fms2["B1"].comment = Comment(
    "평가자가 2명이었고 두 사람 점수를 따로 갖고 있을 때만 입력(선택). 평가자1 점수는 FMS 시트에. "
    "두 평가자 총점의 평균으로 다시 계산한 결과가 자동검산 B-2에 나옵니다.", "양식", width=380, height=200)

# ================================================================== 출석_훈련기록
NSESS = 36
COLS_A = ([("ID", "대상자 ID\n(직접 입력)", "in", 9, "@"), ("group", "집단\n(자동)", "calc", 6, "0")]
          + [(f"s{k:02d}", f"{k}회기\n({(k - 1) // 3 + 1}주)", "in", 5, "0") for k in range(1, NSESS + 1)]
          + [("attend_n", "출석\n회기 수\n(자동)", "calc", 7, "0"),
             ("record_n", "기록된\n회기 수\n(자동)", "calc", 7, "0"),
             ("attend_rate", "출석률\n(실시 회기\n기준, 자동)", "calc", 9, "0.0%"),
             ("hr_main_pct", "본운동 구간\n평균 심박\n(%HRmax)", "in", 9, "0.0"),
             ("hr_n", "심박 기록\n회기 수", "in", 7, "0"),
             ("rpe_mean", "평균 세션\nRPE\n(척도: D6)", "in", 8, "0.0"),
             ("reg_hr_wk", "H-IFIT 제외\n주당 정규훈련\n시간(시간/주)", "in", 10, "0.0"),
             ("reg_rpe", "정규훈련\n평균 강도\n(RPE)", "in", 8, "0.0"),
             ("adverse", "이상반응·부상·특이사항", "in", 30, "@")])
CA = col_map(COLS_A)
HELD_ROW, MIN_ROW = LAST + 2, LAST + 3
s_first, s_last = CA["s01"], CA[f"s{NSESS:02d}"]


def formulas_att(r):
    s = f"{s_first}{r}:{s_last}{r}"
    held = f"${s_first}${HELD_ROW}:${s_last}${HELD_ROW}"
    att = f"SUMPRODUCT(({s}=1)*(({held}=1)+({held}=\"\")))"   # 전체 취소(0) 회기의 출석은 세지 않음
    denom = f"({NSESS}-COUNTIF({held},0))"                      # 35행 빈칸 = 실시
    return {
        "group": 1 if r == EX_ROW else lookup("group", r),
        "attend_n": f'=IF(COUNT({s})=0,"",{att})',
        "record_n": f'=IF(COUNT({s})=0,"",COUNT({s}))',
        "attend_rate": f'=IF(COUNT({s})=0,"",IF({denom}=0,"",{att}/{denom}))',
    }


EX_A = {"ID": "예시", "hr_main_pct": 88.5, "hr_n": 30, "rpe_mean": 16.2, "reg_hr_wk": 6.0, "reg_rpe": 13.0,
        "adverse": "8주차(23회기) 결석: 감기"}
EX_A.update({f"s{k:02d}": (0 if k == 23 else 1) for k in range(1, NSESS + 1)})
layout_data_sheet(ws_att, COLS_A, formulas_att, 58, 3, EX_A)
ws_att[f"{CA['adverse']}{EX_ROW}"].alignment = LEFT_NOWRAP
for row in range(FIRST, LAST + 1):
    ws_att[f"{CA['adverse']}{row}"].alignment = Alignment(horizontal="left", vertical="center")
add_dv(ws_att, f"{s_first}{FIRST}:{s_last}{LAST}", type="list", formula1='"1,0"', allow_blank=True,
       showErrorMessage=True, errorTitle="출석", error="1(출석) 또는 0(결석). 기록이 없으면 빈칸(0을 쓰지 마세요).")
for rr, lab in ((HELD_ROW, "전체 취소된 회기에만 0 (빈칸·1 = 실시)"),
                (MIN_ROW, "그 회기 실제 본운동 시간 (분)")):
    ws_att.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
    c = ws_att.cell(row=rr, column=1, value=lab)
    c.font, c.alignment, c.fill = F_BOLD, WRAP, FILL_SEC
    ws_att.row_dimensions[rr].height = 42
    for k in range(NSESS):
        cc = ws_att.cell(row=rr, column=3 + k)
        cc.fill, cc.font, cc.border, cc.alignment = FILL_IN, F_INPUT, BOX, MID
        cc.number_format = "0"
add_dv(ws_att, f"{s_first}{HELD_ROW}:{s_last}{HELD_ROW}", type="list", formula1='"1,0"', allow_blank=True,
       showErrorMessage=True, errorTitle="회기 실시", error="1(실시) 또는 0(전체 취소).")
ws_att.cell(row=MIN_ROW + 1, column=1, value=(
    "실험군: 회기별 1=출석, 0=결석, 기록 없으면 빈칸. 출석률 = 출석 회기 ÷ 실제 실시 회기. 35행에는 전체 취소된 회기에만 0을 넣고 나머지는 비워 두면 실시로 계산합니다. "
    "통제군은 회기 칸을 비우고 오른쪽 정규훈련 시간·강도만 적어도 됩니다. 실험군의 '정규훈련 시간'은 H-IFIT를 뺀 시간입니다.")
).font = F_NOTE

# ================================================================== 주차별_프로그램
COLS_P = [("week", "주차", "fix", 6, "0"), ("exercises", "실제 실시한 운동 종목 (쉼표로 구분)", "in", 40, "@"),
          ("n_ex", "종목 수", "in", 7, "0"), ("work_s", "종목당\n운동(초)", "in", 9, "0"),
          ("rest_s", "종목 간\n휴식(초)", "in", 9, "0"), ("sets", "세트 수", "in", 7, "0"),
          ("set_rest_s", "세트 간\n휴식(초)", "in", 9, "0"), ("ratio", "운동:휴식\n(자동)", "calc", 9, "@"),
          ("work_min", "순수 운동\n시간(분, 자동)", "calc", 10, "0.0"),
          ("main_min", "본운동 총시간\n추정(분, 자동)", "calc", 11, "0.0"),
          ("diff40", "원고 40분과의\n차이(분, 자동)", "calc", 11, "+0.0;-0.0;0.0"),
          ("target_hr", "목표 강도\n(%HRmax)", "in", 9, "@"), ("actual_hr", "실제 평균\n심박(%HRmax)", "in", 10, "0.0"),
          ("target_rpe", "목표 RPE", "in", 8, "@"), ("note", "진행(점증) 기준·비고", "in", 30, "@")]
CP = col_map(COLS_P)
ws_prog["A1"] = "주차별 실제 운동 구성 (프로그램 재현성)"
ws_prog["A1"].font = F_TITLE
ws_prog["A2"] = ("주차마다 실제로 한 구성을 적습니다(회기마다 달랐으면 대표 구성 + 비고). 본운동 총시간 추정 = "
                 "종목 수×세트×(운동+휴식) + (세트−1)×세트 간 휴식. 전환 시간은 빠져 있으니 연구정보 D9 답과 함께 봅니다.")
ws_prog["A2"].font, ws_prog["A2"].alignment = F_NOTE, WRAP
ws_prog.merge_cells("A2:O2")
ws_prog.row_dimensions[2].height = 30
PH = 4
for i, (name, lab, kind, w, fmt) in enumerate(COLS_P, start=1):
    hdr_cell(ws_prog, PH - 1, i, name, font=F_VAR).alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)
    c = ws_prog.cell(row=PH, column=i, value=lab)
    c.font, c.alignment, c.border = F_BOLD, CENTER, BOX
    c.fill = FILL_CALC if kind == "calc" else FILL_SEC
    ws_prog.column_dimensions[L(i)].width = w
ws_prog.row_dimensions[PH].height = 45
EX_P = {"exercises": "스쿼트 점프, 점핑 런지, 버피, 플랭크, 셔틀런, 교대 돌려차기", "n_ex": 6, "work_s": 30,
        "rest_s": 30, "sets": 3, "set_rest_s": 120, "target_hr": "85–90%", "actual_hr": 87.2,
        "target_rpe": "15–16", "note": "적응기: 1:1"}
for idx in range(13):
    row = PH + 1 + idx
    for i, (name, lab, kind, w, fmt) in enumerate(COLS_P, start=1):
        c = ws_prog.cell(row=row, column=i)
        c.border, c.number_format = BOX, fmt
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=name in ("exercises", "note"))
        n_, w_, r_, s_, sr_ = (f"{CP[k]}{row}" for k in ("n_ex", "work_s", "rest_s", "sets", "set_rest_s"))
        if name == "week":
            c.value = "예시" if idx == 0 else idx
            c.font, c.fill = (F_EX, FILL_EX) if idx == 0 else (F_BOLD, FILL_SEC)
        elif kind == "calc":
            c.fill, c.font = FILL_CALC, (F_EX if idx == 0 else F_BASE)
            c.value = {
                "ratio": f'=IF(AND(ISNUMBER({w_}),ISNUMBER({r_})),IF({r_}>0,TEXT({w_}/{r_},"0.0#")&" : 1",""),"")',
                "work_min": f'=IF(COUNT({n_},{w_},{s_})=3,{n_}*{s_}*{w_}/60,"")',
                "main_min": (f'=IF(COUNT({n_},{w_},{r_},{s_})=4,({n_}*{s_}*({w_}+{r_})'
                             f'+MAX({s_}-1,0)*IF(ISNUMBER({sr_}),{sr_},0))/60,"")'),
                "diff40": f'=IF(ISNUMBER({CP["main_min"]}{row}),{CP["main_min"]}{row}-40,"")',
            }[name]
        elif idx == 0:
            c.value, c.fill, c.font = EX_P.get(name), FILL_EX, F_EX
        else:
            c.fill, c.font = FILL_IN, F_INPUT
ws_prog.freeze_panes = ws_prog.cell(row=PH + 1, column=2)

# ================================================================== 연구정보
ws_info["A1"] = "연구 정보 (원고 '연구방법' 작성용)"
ws_info["A1"].font = F_TITLE
ws_info["A2"] = ("노란 칸에 답하고, E열에 근거(문서·기록 / 기억 / 모름)를 고르세요. 모르면 '모름', 하지 않았으면 '미실시'라고 "
                 "적어 주세요. 사실과 다른 답보다 '모름'이 훨씬 낫습니다. '기억'으로 표시한 답은 원고에 조심스럽게 쓰거나 "
                 "투고 전에 기록으로 확인합니다.")
ws_info["A2"].font, ws_info["A2"].alignment = F_NOTE, WRAP
ws_info.merge_cells("A2:E2")
ws_info.row_dimensions[2].height = 42
INFO = [
    ("A. 설계·배정", None, None),
    ("A1", "무작위 배정을 했는가 (예 / 아니오)", "예=무작위 대조시험(CONSORT), 아니오=비무작위 대조연구(TREND)"),
    ("A2", "배정 단위 (개인 / 팀·학교)", "팀 단위면 군집 구조를 분석·기술해야 함"),
    ("A3", "무작위라면 방법 (난수표·제비뽑기·블록/층화 여부), 배정한 사람, 배정 은폐 방법", "CONSORT 필수 항목"),
    ("A4", "무작위가 아니라면 배정 기준 (자원, 지도자 지정, 소속 팀 등)", "선택편향 논의"),
    ("A5", "측정자 눈가림 여부 (VO2max·FMS 측정자가 집단을 알았는가)", "측정 편향 대응"),
    ("A6", "연구계획서·분석계획 문서가 있는가, 작성 시점 (자료 수집 전/후), 일차·이차 결과변수와 가설", "결과를 보고 가설을 세웠다는 의심 방지"),
    ("A7", "임상연구 등록 여부, 등록처 (CRIS 등), 등록번호, 등록일 (미등록이면 '미등록')", "CONSORT·TREND 보고 항목"),
    ("B. 대상자", None, None),
    ("B1", "소속 기관·팀 (익명 가능: 예) 경기도 소재 A대학 태권도부)", "대상 기술"),
    ("B2", "모집 기간, 중재 시작일과 종료일", "중재 시점·시합 일정 확인"),
    ("B3", "선정 기준 (예: 경력 5년 이상, 최근 6개월 부상 없음)", "필수"),
    ("B4", "제외 기준", "필수"),
    ("B5", "표본수 산출 근거 (G*Power: 효과크기 f, α, 검정력, 산출 인원). 사전에 하지 않았으면 '사전 산출 없음'", "심사 단골 지적. 사후에 맞춰 만들지 않음"),
    ("B6", "시즌 구분 (시합기 / 준비기 / 비시즌)과 중재 기간 중 출전한 대회", "교란변수"),
    ("C. 연구윤리", None, None),
    ("C1", "IRB 심의 상태 (승인 / 면제 / 심의 받지 않음)와 심의 기관명", "사실대로 기재. 사후 승인이면 사후 승인으로 표기"),
    ("C2", "IRB 승인번호, 승인일. 승인일이 첫 사전 측정일보다 앞서는가", "본문 기재 사항"),
    ("C3", "서면 동의 취득 여부 (미성년자 포함 시 보호자 동의)", "필수"),
    ("C4", "개인정보 보호 방법 (코드화, 자료 보관자·장소·기간)", "연구윤리 기술"),
    ("C5", "중재 중 안전관리 (사전 건강 확인, 현장 인원, 운동 중단 기준)", "고강도 중재의 안전성 기술"),
    ("D. 중재·비교조건", None, None),
    ("D1", "실험군은 정규훈련에 H-IFIT를 '추가'했는가, 정규훈련 일부를 H-IFIT로 '대체'했는가", "핵심: 두 집단 총 운동량 차이"),
    ("D2", "통제군 정규훈련 내용 (겨루기·기술·체력 등 구성)", "비교조건 기술"),
    ("D3", "정규훈련 주당 횟수·시간·강도(RPE 등): 두 집단 각각, 중재 기간 중", "총 훈련량 비교"),
    ("D4", "H-IFIT 지도자 (자격·경력), 장소, 장비", "재현성"),
    ("D5", "HRmax 산출 방법 (검사 실측 / 220−연령 / 기타), 심박 측정 장비 모델, 심박 기록 회기", "강도 보고"),
    ("D6", "RPE 척도 (Borg 6–20 / CR-10)와 수집 시점 (회기 종료 후 몇 분)", "강도 보고"),
    ("D7", "주차별 진행(점증) 기준 (예: 4주마다 운동:휴식 비율 변경)", "주차별_프로그램 시트와 함께"),
    ("D8", "두 집단의 중재 기간 중 부상·이상반응 (없으면 '없음, 기록 있음' / '기록 없음' 구분)", "위해 보고"),
    ("D9", "'본운동 40분'에 휴식·세트 간 휴식·전환 시간이 포함되는가", "포스터 범위(6–10종목·3–4세트)와 40분의 정합성"),
    ("D10", "12개 운동 중 회기별 6–10개를 고른 기준과 실시 순서", "재현성"),
    ("D11", "수면·영양·개인 추가훈련 통제 여부 (두 집단)", "교란변수"),
    ("E. VO2max 측정", None, None),
    ("E1", "트레드밀 제조사·모델", "필수"),
    ("E2", "호흡가스 분석기 제조사·모델 (직접 측정인지, 추정식인지)", "직접 측정이 아니면 '추정 VO2max'로 표기"),
    ("E3", "가스분석기 보정 절차 (매 검사 전 가스·유량 보정 등)", "측정 신뢰도"),
    ("E4", "부하 프로토콜 (Bruce / 수정 Bruce / 자체: 단계별 속도·경사·시간)", "필수"),
    ("E5", "종료 기준과 최대치 판정 기준 (예: RER ≥ 1.10, 최고심박 ≥ 연령예측 95%, VO2 plateau, RPE ≥ 17 중 2개 이상)", "VO2max / VO2peak 명칭 결정. 대상자 시트의 '충족 기준 수'는 이 예시 기준으로 계산"),
    ("E6", "VO2max 값 산출 방식 (30초 평균 최고값 / 15초 / breath-by-breath)", "방법 기술"),
    ("E7", "VO2max·FMS 각각의 측정 시간대와 사전 금지사항 (24시간 고강도 운동·카페인 금지 등)", "표준화"),
    ("E8", "사후 측정은 마지막 훈련 후 며칠(몇 시간) 뒤였는가 (VO2max·FMS 각각)", "피로 영향"),
    ("F. FMS 측정", None, None),
    ("F1", "평가자 수와 FMS 자격(Level 1 등)·경력", "필수"),
    ("F2", "사전·사후 평가자가 같은 사람인가", "측정 일관성"),
    ("F3", "실시간 판정인가, 영상 촬영 후 판정인가 (촬영 각도)", "신뢰도"),
    ("F4", "각 항목 시도 횟수와 채점 규칙 (3회 중 최고점 등)", "매뉴얼 준수"),
    ("F5", "평가자가 2명 이상이면 최종 점수 결정 방식 (합의 / 평균 / 1명 점수 사용)", "총점이 소수가 되는 원인 확인"),
    ("F6", "평가자 간·평가자 내 신뢰도 산출 여부와 값 (ICC, kappa)", "심사 단골 지적. 두 평가자 점수가 있으면 FMS_평가자2 시트로 계산 가능"),
    ("F7", "Clearing test(어깨·척추 신전·척추 굴곡) 실시 여부", "통증이면 해당 항목 0점"),
    ("F8", "FMS 측정 순서 (VO2max 검사 전/후, 같은 날/다른 날)", "피로 영향"),
    ("F9", "좌우 점수 처리 방식 (낮은 쪽 점수 / 좌우 평균 / 좌우 합산 / 기타)과 총점 만점 (21점 / 기타)", "이 양식은 '낮은 쪽, 21점 만점'(Cook 2014)으로 계산. 다르면 포스터와 같은 기준이 아님"),
    ("F10", "적용한 FMS 매뉴얼·판본 (Cook 2014 / FMS Level 1 교재 등)", "방법 기술"),
    ("F11", "선수들이 연구 전에 FMS를 받아 보거나 연습한 적이 있는가", "학습효과 교란"),
    ("G. 분석", None, None),
    ("G1", "포스터 수치를 낸 SPSS 파일(.sav)과 출력(.spv)이 있는가", "수치 재현"),
    ("G2", "포스터 분석에 실제 쓰인 인원: VO2max (실험 n, 통제 n), FMS (실험 n, 통제 n)", "결과별 분석 인원"),
    ("G3", "포스터 F값을 낸 분석 (SPSS 메뉴: 일반선형모형-반복측정 / 공분산분석 등)과 공변량", "자동검산의 F와 다를 때 원인 확인"),
    ("G4", "분석 단위 (개인 / 팀)와 같은 팀 선수 포함 여부", "관측 독립성"),
    ("G5", ".spv가 없으면 출력에 적힌 자유도·F·t·p 값을 그대로 옮겨 적기", "자유도 보고"),
    ("G6", "결측치 처리 방법 (완전사례 분석 / 기타)", "방법 기술"),
    ("G7", "정규성 검정 실시 여부와 결과 (Shapiro-Wilk 등)", "모수/비모수 선택 근거"),
    ("G8", "FMS 하위항목 검정: 방법 (대응 t / Wilcoxon / 기타), 집단 내 전후 비교인가 집단×시기인가, 다중비교 보정, 항목별 검정값·p (집단별)", "원고 표 5 해석"),
    ("G9", "사후검정 실시 여부", "방법 기술"),
    ("G10", "포스터 분석을 수행한 사람", "자료 추적"),
    ("G11", "이 양식에 입력한 값의 출처: VO2max (가스분석기 출력물 / 기록지 / 기존 SPSS), FMS (채점지 / 영상 / 기존 SPSS)", "원자료 검증인지 SPSS 재집계인지 구분"),
    ("H. 투고 정보", None, None),
    ("H1", "포스터 발표: 공식 행사명, 발표일, 장소, 초록집 수록 여부·쪽수", "선행 발표 고지"),
    ("H2", "포스터 'Research Background'의 'H-IFIT가 Deep Squat를 유의하게 개선' 문장 출처 (본 연구 결과 / 선행연구)", "원고에서 출처 불명으로 지적됨"),
    ("H3", "저자 전원의 국문·영문 성명, 국문·영문 소속, 직위, 이메일, 최종 저자 순서", "원고에 '확인 필요'로 남아 있음"),
    ("H4", "교신저자 (Ki-Seung Ham)의 국문 성명, 소속, 직위, 이메일", "영문 표기에서 국문 이름을 추정하지 않음"),
    ("H5", "연구비 지원 여부, 이해상충", "필수"),
    ("H6", "저자별 기여 (연구설계, 자료수집, 분석, 집필)", "학술지 요구 시"),
    ("H7", "연구 전 과정에서 쓴 AI 도구명·버전과 사용 단계 (문헌 검색, 원고 작성·번역, 통계 검산, 자료 입력 양식 작성 등), 저자 검증 여부", "국기원 연구윤리지침상 공개"),
    ("H8", "자료 이용 가능성 (요청 시 제공 / 비공개와 사유)", "원고 기재 사항"),
    ("H9", "학위논문의 일부인가, 포스터 외 다른 학술지·학술대회에 발표·게재했는가", "중복게재 고지"),
]
IH = 4
for j, h in enumerate(["번호", "확인 항목", "입력 (답변)", "왜 필요한가", "근거"], start=1):
    hdr_cell(ws_info, IH, j, h)
row = IH + 1
for code, text, why in INFO:
    if text is None:
        ws_info.cell(row=row, column=1, value=code).font = F_BOLD
        for j in range(1, 6):
            ws_info.cell(row=row, column=j).fill = FILL_SEC
            ws_info.cell(row=row, column=j).border = BOX
        ws_info.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    else:
        for j, v in enumerate([code, text, None, why, None], start=1):
            c = ws_info.cell(row=row, column=j, value=v)
            c.border, c.alignment = BOX, WRAP
            c.font = F_NOTE if j == 4 else (F_INPUT if j in (3, 5) else F_BASE)
            if j in (3, 5):
                c.fill = FILL_IN
        ws_info.cell(row=row, column=1).alignment = CENTER
        ws_info.cell(row=row, column=5).alignment = CENTER
        ws_info.row_dimensions[row].height = 48
    row += 1
INFO_LAST = row - 1
add_dv(ws_info, f"E{IH + 1}:E{INFO_LAST}", type="list", formula1='"문서·기록,기억,모름"', allow_blank=True,
       showErrorMessage=True, errorStyle="information", errorTitle="근거", error="문서·기록 / 기억 / 모름 중 선택")
for col, w in zip("ABCDE", (6, 56, 50, 36, 10)):
    ws_info.column_dimensions[col].width = w
ws_info.freeze_panes = f"A{IH + 1}"

# ================================================================== 대상자_흐름
ws_flow["A1"] = "대상자 흐름 (CONSORT/TREND 흐름도)"
ws_flow["A1"].font = F_TITLE
ws_flow["A2"] = ("각 선수는 단계마다 한 번만 셉니다. '중재 중단'은 중재를 그만뒀지만 사후 측정은 받은 경우, "
                 "'추적 소실'은 사후 측정을 받지 못한 경우입니다. 중단했지만 사후 측정을 받은 선수는 분석에 포함될 수 있습니다.")
ws_flow["A2"].font, ws_flow["A2"].alignment = F_NOTE, WRAP
ws_flow.merge_cells("A2:F2")
ws_flow.row_dimensions[2].height = 42
for j, h in enumerate(["단계", "실험군", "통제군", "합계", "사유·설명 (직접 입력)", "점검"], start=1):
    hdr_cell(ws_flow, 4, j, h)
grp_s = f"{q(S_SUBJ)}!${CS['group']}${FIRST}:${CS['group']}${LAST}"
FLOW = [
    ("적격성 평가 인원", "total"),
    ("배정 전 제외: 선정기준 미충족", "total"),
    ("배정 전 제외: 참여 거부", "total"),
    ("배정 전 제외: 기타", "total"),
    ("배정 인원", "in"),
    ("배정된 중재를 실제로 받음", "in"),
    ("중재 중단 (사후 측정은 받음)", "in"),
    ("추적 소실 (사후 측정 못 받음)", "in"),
    ("VO2max 분석 인원 (직접 입력)", "in"),
    ("FMS 분석 인원 (직접 입력)", "in"),
    ("VO2max 분석 인원 (자동: 자동검산 A)", "auto"),
    ("FMS 분석 인원 (자동: 자동검산 B)", "auto"),
    ("대상자 시트 입력 인원 (자동)", "auto_all"),
]
FR = {}
for k, (lab, kind) in enumerate(FLOW):
    r = 5 + k
    FR[lab] = r
    ws_flow.cell(row=r, column=1, value=lab)
    for j in range(1, 7):
        c = ws_flow.cell(row=r, column=j)
        c.border, c.font = BOX, F_BASE
        c.alignment = WRAP if j in (1, 5, 6) else CENTER
    if kind == "total":
        for j in (2, 3):
            ws_flow.cell(row=r, column=j, value="—").fill = FILL_CALC
        ws_flow.cell(row=r, column=4).fill, ws_flow.cell(row=r, column=4).font = FILL_IN, F_INPUT
    elif kind == "in":
        for j in (2, 3):
            ws_flow.cell(row=r, column=j).fill, ws_flow.cell(row=r, column=j).font = FILL_IN, F_INPUT
        ws_flow.cell(row=r, column=4, value=f'=IF(COUNT(B{r}:C{r})=0,"",SUM(B{r}:C{r}))').fill = FILL_CALC
    else:
        for j in (2, 3, 4):
            ws_flow.cell(row=r, column=j).fill = FILL_CALC
        if kind == "auto_all":
            ws_flow.cell(row=r, column=2, value=f"=COUNTIF({grp_s},1)")
            ws_flow.cell(row=r, column=3, value=f"=COUNTIF({grp_s},2)")
        ws_flow.cell(row=r, column=4, value=f"=B{r}+C{r}")
    ws_flow.cell(row=r, column=5).fill, ws_flow.cell(row=r, column=5).font = FILL_IN, F_INPUT
    ws_flow.row_dimensions[r].height = 28
for col, w in zip("ABCDEF", (36, 9, 9, 9, 40, 34)):
    ws_flow.column_dimensions[col].width = w

# ================================================================== 정의된 이름
NAMES = {
    "grp": (S_SUBJ, CS["group"]), "inc": (S_SUBJ, CS["include"]),
    "vo2_pre": (S_SUBJ, CS["vo2_pre"]), "vo2_post": (S_SUBJ, CS["vo2_post"]), "d_vo2": (S_SUBJ, CS["d_vo2"]),
    "wt_pre": (S_SUBJ, CS["wt_pre"]), "wt_post": (S_SUBJ, CS["wt_post"]), "d_wt": (S_SUBJ, CS["d_wt"]),
    "age": (S_SUBJ, CS["age"]), "height": (S_SUBJ, CS["height"]), "career": (S_SUBJ, CS["career"]),
    "sex": (S_SUBJ, CS["sex"]),
    "grp_f": (S_FMS, CF["group"]), "inc_f": (S_FMS, CF["include"]),
    "fms_pre": (S_FMS, CF["fms_pre"]), "fms_post": (S_FMS, CF["fms_post"]), "d_fms": (S_FMS, CF["d_fms"]),
    "fms_avg_pre": (S_FMS, CF["avg_pre"]), "fms_avg_post": (S_FMS, CF["avg_post"]), "d_fms_avg": (S_FMS, CF["d_avg"]),
    "asym_pre": (S_FMS, CF["asym_pre"]), "asym_post": (S_FMS, CF["asym_post"]),
}
for code, *_ in ITEMS:
    for t in ("pre", "post"):
        NAMES[f"{code}_{t}_s"] = (S_FMS, CF[f"{code}s_{t}"])
for n, (sh, col) in NAMES.items():
    wb.defined_names[n] = DefinedName(n, attr_text=f"{q(sh)}!${col}${FIRST}:${col}${LAST}")

# ================================================================== 자동검산
wc = ws_calc
wc["A1"] = "자동 검산 (입력 자료 → 기술통계·검정값 → 포스터 보고값과의 차이 기록)"
wc["A1"].font = F_TITLE
wc["A2"] = ("집단 1/2, 분석포함=1, 사전·사후가 모두 있는 대상자만 씁니다(완전사례). 집단×시기 상호작용 F는 두 시점 혼합 "
            "반복측정 ANOVA의 상호작용과 수학적으로 같은 '변화량의 독립표본 t²(등분산 가정)'입니다. FMS 총점은 "
            "'좌우 중 낮은 점수, clearing 통증=0, 21점 만점'(Cook 2014) 규칙으로 계산합니다(연구정보 F9 참고). "
            "하위항목 비모수 검정·ICC 등 최종 분석은 자료를 받은 뒤 따로 합니다.")
wc["A2"].font, wc["A2"].alignment = F_NOTE, WRAP
wc.merge_cells("A2:N2")
wc.row_dimensions[2].height = 56


def arr(ref, formula):
    wc[ref] = ArrayFormula(ref, formula)


def style_row(r, ncol, fmts):
    for j in range(1, ncol + 1):
        c = wc.cell(row=r, column=j)
        c.border, c.alignment = BOX, CENTER
        if j > 1:
            c.font, c.fill = F_BASE, FILL_CALC
        c.number_format = fmts.get(j, "0.00")


def block_desc(top, title, g, inc, pre, post, d, unit, extra=""):
    wc.cell(row=top, column=1, value=title).font = F_SEC
    heads = ["집단", "n", "사전 M", "사전 SD", "사후 M", "사후 SD", "변화 M\n(사후−사전)", "변화 SD",
             "대응 t\n(사후−사전)", "df", "p (양측)", "변화 95% CI\n하한", "변화 95% CI\n상한", "효과크기 dz"]
    for j, h in enumerate(heads, start=1):
        hdr_cell(wc, top + 1, j, h)
    wc.row_dimensions[top + 1].height = 32
    rows = []
    for k, (gv, gname) in enumerate(((1, "실험군"), (2, "통제군"))):
        r = top + 2 + k
        rows.append(r)
        wc.cell(row=r, column=1, value=gname).font = F_BOLD
        cnd = f"({g}={gv})*({inc}=1)*ISNUMBER({d}){extra}"
        wc[f"B{r}"] = f"=SUMPRODUCT({cnd})"
        arr(f"C{r}", f'=IF(B{r}<1,"",AVERAGE(IF({cnd},{pre})))')
        arr(f"D{r}", f'=IF(B{r}<2,"",STDEV(IF({cnd},{pre})))')
        arr(f"E{r}", f'=IF(B{r}<1,"",AVERAGE(IF({cnd},{post})))')
        arr(f"F{r}", f'=IF(B{r}<2,"",STDEV(IF({cnd},{post})))')
        arr(f"G{r}", f'=IF(B{r}<1,"",AVERAGE(IF({cnd},{d})))')
        arr(f"H{r}", f'=IF(B{r}<2,"",STDEV(IF({cnd},{d})))')
        wc[f"I{r}"] = f'=IF(OR(B{r}<2,H{r}=""),"",IF(ROUND(H{r},10)=0,"",G{r}/(H{r}/SQRT(B{r}))))'
        wc[f"J{r}"] = f'=IF(B{r}<2,"",B{r}-1)'
        wc[f"K{r}"] = f'=IF(I{r}="","",TDIST(ABS(I{r}),J{r},2))'
        wc[f"L{r}"] = f'=IF(I{r}="","",G{r}-TINV(0.05,J{r})*H{r}/SQRT(B{r}))'
        wc[f"M{r}"] = f'=IF(I{r}="","",G{r}+TINV(0.05,J{r})*H{r}/SQRT(B{r}))'
        wc[f"N{r}"] = f'=IF(I{r}="","",G{r}/H{r})'
        style_row(r, 14, {2: "0", 9: "0.000", 10: "0", 11: "0.0000", 14: "0.00"})
    wc.cell(row=top + 4, column=1, value=(
        f"단위: {unit}. t는 사후−사전 기준(양수=증가). SPSS는 보통 사전−사후로 계산해 부호가 반대입니다.")).font = F_NOTE
    return rows


def block_between(top, title, e, c_, mcol, sdcol, label, note=None):
    wc.cell(row=top, column=1, value=title).font = F_SEC
    heads = ["비교", "n 실험", "n 통제", f"{label} 차\n(실험−통제)", "합동 SD", "독립 t", "df", "F(1, df)",
             "p", "부분 η²", "차이 95% CI\n하한", "차이 95% CI\n상한"]
    for j, h in enumerate(heads, start=1):
        hdr_cell(wc, top + 1, j, h)
    wc.row_dimensions[top + 1].height = 32
    r = top + 2
    wc.cell(row=r, column=1, value="실험 vs 통제").font = F_BOLD
    wc[f"B{r}"], wc[f"C{r}"] = f"=B{e}", f"=B{c_}"
    wc[f"D{r}"] = f'=IF(OR(B{r}<2,C{r}<2),"",{mcol}{e}-{mcol}{c_})'
    wc[f"E{r}"] = f'=IF(D{r}="","",SQRT(((B{r}-1)*{sdcol}{e}^2+(C{r}-1)*{sdcol}{c_}^2)/(B{r}+C{r}-2)))'
    wc[f"F{r}"] = f'=IF(D{r}="","",IF(ROUND(E{r},10)=0,"",D{r}/(E{r}*SQRT(1/B{r}+1/C{r}))))'
    wc[f"G{r}"] = f'=IF(F{r}="","",B{r}+C{r}-2)'
    wc[f"H{r}"] = f'=IF(F{r}="","",F{r}^2)'
    wc[f"I{r}"] = f'=IF(F{r}="","",TDIST(ABS(F{r}),G{r},2))'
    wc[f"J{r}"] = f'=IF(F{r}="","",H{r}/(H{r}+G{r}))'
    wc[f"K{r}"] = f'=IF(F{r}="","",D{r}-TINV(0.05,G{r})*E{r}*SQRT(1/B{r}+1/C{r}))'
    wc[f"L{r}"] = f'=IF(F{r}="","",D{r}+TINV(0.05,G{r})*E{r}*SQRT(1/B{r}+1/C{r}))'
    style_row(r, 12, {2: "0", 3: "0", 6: "0.000", 7: "0", 8: "0.000", 9: "0.0000", 10: "0.000"})
    if note:
        wc.cell(row=r + 1, column=1, value=note).font = F_NOTE
    return r


cur = 4
VE, VC = block_desc(cur, "A. VO2max 집단별 사전·사후 (mL/kg/min)", "grp", "inc", "vo2_pre", "vo2_post", "d_vo2", "mL/kg/min"); cur += 6
FE, FC = block_desc(cur, "B. FMS 총점 집단별 사전·사후 (점) — FMS 시트(평가자1, 하위항목이 없으면 SPSS총점)", "grp_f", "inc_f", "fms_pre", "fms_post", "d_fms", "점(21점 만점)"); cur += 6
AE, AC = block_desc(cur, "B-2. FMS 총점: 두 평가자 총점의 평균 (FMS_평가자2 시트 입력 시)", "grp_f", "inc_f", "fms_avg_pre", "fms_avg_post", "d_fms_avg", "점(두 평가자 평균)"); cur += 6
WE, WC = block_desc(cur, "C. 체중 집단별 사전·사후 (kg) — A와 같은 대상자, 상대 VO2max 해석용", "grp", "inc", "wt_pre", "wt_post", "d_wt", "kg", extra="*ISNUMBER(d_vo2)"); cur += 6
VI = block_between(cur, "D. 집단×시기 상호작용: VO2max (변화량의 집단 비교)", VE, VC, "G", "H", "변화 M"); cur += 4
FI = block_between(cur, "E. 집단×시기 상호작용: FMS 총점", FE, FC, "G", "H", "변화 M"); cur += 4
AI = block_between(cur, "E-2. 집단×시기 상호작용: FMS 두 평가자 평균 총점", AE, AC, "G", "H", "변화 M"); cur += 4
NOTE_BASE = "사전 동질성은 독립 t와 p, 효과크기만 보세요(F·η² 열은 해당 없음). p > .05가 '동질성 입증'은 아닙니다."
VB = block_between(cur, "F. 사전 동질성: VO2max 사전값", VE, VC, "C", "D", "사전 M", NOTE_BASE); cur += 4
FB = block_between(cur, "G. 사전 동질성: FMS 사전 총점", FE, FC, "C", "D", "사전 M", NOTE_BASE); cur += 5

# H. 포스터 대조
H_TOP = cur
wc.cell(row=H_TOP, column=1, value="H. 포스터 보고값과 입력 자료 재계산값의 차이 (기록용)").font = F_SEC
wc.cell(row=H_TOP + 1, column=1, value=(
    "'다름'은 고칠 대상이 아니라 그 자체로 보고할 결과입니다. 원 기록지·측정 출력물의 값을 그대로 입력하고, 포스터 값에 "
    "맞추려고 원자료를 고치거나 대상자를 제외(분석포함=0)하지 마세요. 입력 실수를 고쳤다면 '수정기록' 시트에 남깁니다. "
    "포스터 값 출처: 사용자 제공 원고(3차) 표 3·4. t는 SPSS 관례(사전−사후)로 바꿔 부호까지 비교합니다.")).font = F_RED
wc.merge_cells(start_row=H_TOP + 1, start_column=1, end_row=H_TOP + 1, end_column=9)
wc.cell(row=H_TOP + 1, column=1).alignment = WRAP
wc.row_dimensions[H_TOP + 1].height = 58
heads = ["항목", "포스터 값", "입력 자료\n재계산", "차이", "판정", "t 부호",
         "두 평가자 평균\n기준 재계산", "차이", "판정"]
for j, h in enumerate(heads, start=1):
    hdr_cell(wc, H_TOP + 2, j, h)
wc.row_dimensions[H_TOP + 2].height = 32
# 라벨, 포스터 값, 재계산 셀, 자릿수, n 셀, 평균기준 셀, 평균기준 n 셀, 종류(n/v/t)
CMP = [
    ("VO2max 실험군 n", 13, f"B{VE}", 0, None, None, None, "n"),
    ("VO2max 통제군 n", 12, f"B{VC}", 0, None, None, None, "n"),
    ("VO2max 실험군 사전 M", 48.31, f"C{VE}", 2, f"B{VE}", None, None, "v"),
    ("VO2max 실험군 사전 SD", 4.12, f"D{VE}", 2, f"B{VE}", None, None, "v"),
    ("VO2max 실험군 사후 M", 52.38, f"E{VE}", 2, f"B{VE}", None, None, "v"),
    ("VO2max 실험군 사후 SD", 4.05, f"F{VE}", 2, f"B{VE}", None, None, "v"),
    ("VO2max 실험군 t", -6.184, f"I{VE}", 3, f"B{VE}", None, None, "t"),
    ("VO2max 통제군 t", -1.207, f"I{VC}", 3, f"B{VC}", None, None, "t"),
    ("VO2max 통제군 p", 0.253, f"K{VC}", 3, f"B{VC}", None, None, "v"),
    ("VO2max 상호작용 F", 18.742, f"H{VI}", 3, f"B{VE}", None, None, "v"),
    ("FMS 실험군 n", 13, f"B{FE}", 0, None, f"B{AE}", None, "n"),
    ("FMS 통제군 n", 12, f"B{FC}", 0, None, f"B{AC}", None, "n"),
    ("FMS 실험군 사전 M", 14.92, f"C{FE}", 2, f"B{FE}", f"C{AE}", f"B{AE}", "v"),
    ("FMS 실험군 사전 SD", 1.44, f"D{FE}", 2, f"B{FE}", f"D{AE}", f"B{AE}", "v"),
    ("FMS 실험군 사후 M", 17.08, f"E{FE}", 2, f"B{FE}", f"E{AE}", f"B{AE}", "v"),
    ("FMS 실험군 사후 SD", 1.32, f"F{FE}", 2, f"B{FE}", f"F{AE}", f"B{AE}", "v"),
    ("FMS 실험군 t", -7.368, f"I{FE}", 3, f"B{FE}", f"I{AE}", f"B{AE}", "t"),
    ("FMS 통제군 t", -1.483, f"I{FC}", 3, f"B{FC}", f"I{AC}", f"B{AC}", "t"),
    ("FMS 통제군 p", 0.166, f"K{FC}", 3, f"B{FC}", f"K{AC}", f"B{AC}", "v"),
    ("FMS 상호작용 F", 21.396, f"H{FI}", 3, f"B{FE}", f"H{AI}", f"B{AE}", "v"),
]


def verdict(val, poster, dig, nref, kind):
    if kind == "n":
        return f'=IF(NOT(ISNUMBER({val})),"",IF({val}=0,"자료 입력 전",IF({val}={poster},"같음","다름")))'
    pending = f'IF({nref}=0,"자료 입력 전","계산 불가(n<2 또는 SD=0)")'
    return (f'=IF(NOT(ISNUMBER({val})),{pending},IF(ROUND({val},{dig})=ROUND({poster},{dig}),"같음",'
            f'IF(ABS({val}-{poster})<=1.5*10^(-{dig}),"끝자리 1 차이","다름")))')


for k, (lab, pv, expr, dig, nref, aexpr, anref, kind) in enumerate(CMP):
    r = H_TOP + 3 + k
    wc.cell(row=r, column=1, value=lab).font = F_BASE
    wc.cell(row=r, column=2, value=pv).font = F_INPUT
    wc[f"C{r}"] = f'=IF(ISNUMBER({expr}),{"-" if kind == "t" else ""}{expr},"")'
    wc[f"D{r}"] = f'=IF(ISNUMBER(C{r}),C{r}-B{r},"")'
    wc[f"E{r}"] = verdict(f"C{r}", f"B{r}", dig, nref, kind)
    if kind == "t":
        wc[f"F{r}"] = f'=IF(ISNUMBER(C{r}),IF(SIGN(C{r})=SIGN(B{r}),"같음","반대"),"")'
    if aexpr:
        wc[f"G{r}"] = f'=IF(ISNUMBER({aexpr}),{"-" if kind == "t" else ""}{aexpr},"")'
        wc[f"H{r}"] = f'=IF(ISNUMBER(G{r}),G{r}-B{r},"")'
        wc[f"I{r}"] = verdict(f"G{r}", f"B{r}", dig, anref, kind)
    fmt = "0" if kind == "n" else ("0.00" if dig == 2 else "0.000")
    for j in range(1, 10):
        cc = wc.cell(row=r, column=j)
        cc.border = BOX
        if j > 1:
            cc.alignment = CENTER
        if j >= 3:
            cc.fill = FILL_CALC
        if j in (2, 3, 4, 7, 8):
            cc.number_format = fmt
H_END = H_TOP + 2 + len(CMP)
for col in ("E", "I"):
    wc.conditional_formatting.add(f"{col}{H_TOP + 3}:{col}{H_END}",
                                  FormulaRule(formula=[f'{col}{H_TOP + 3}="같음"'], fill=FILL_OK))
    wc.conditional_formatting.add(f"{col}{H_TOP + 3}:{col}{H_END}",
                                  FormulaRule(formula=[f'{col}{H_TOP + 3}="다름"'], fill=FILL_DIFF))
wc.cell(row=H_END + 1, column=1, value=(
    "포스터 값(파란 글씨)은 원고에 옮겨진 숫자 그대로이며 수정하지 않습니다. FMS를 SPSS총점으로 대체한 값이 있으면(J 참고) "
    "FMS 비교는 원자료 검증이 아니라 기존 SPSS 값의 재집계입니다.")).font = F_NOTE

# I. 하위항목
cur = H_END + 3
wc.cell(row=cur, column=1, value="I. FMS 하위항목 점수 (B와 같은 대상자; SPSS총점만 있는 행은 하위항목 없음)").font = F_SEC
heads = ["항목", "실험 n", "실험 사전 M", "실험 사후 M", "실험 변화", "통제 n", "통제 사전 M", "통제 사후 M", "통제 변화"]
for j, h in enumerate(heads, start=1):
    hdr_cell(wc, cur + 1, j, h)
SUB = [(code, f"{code} ({name})") for code, name, *_ in ITEMS] + [("asym", "좌우 비대칭 항목 수")]
for k, (code, lab) in enumerate(SUB):
    r = cur + 2 + k
    wc.cell(row=r, column=1, value=lab).font = F_BASE
    pre_n = "asym_pre" if code == "asym" else f"{code}_pre_s"
    post_n = "asym_post" if code == "asym" else f"{code}_post_s"
    for gv, base in ((1, 2), (2, 6)):
        cnd = f"(grp_f={gv})*(inc_f=1)*ISNUMBER(d_fms)*ISNUMBER({pre_n})*ISNUMBER({post_n})"
        nref = f"{L(base)}{r}"
        wc[nref] = f"=SUMPRODUCT({cnd})"
        arr(f"{L(base + 1)}{r}", f'=IF({nref}<1,"",AVERAGE(IF({cnd},{pre_n})))')
        arr(f"{L(base + 2)}{r}", f'=IF({nref}<1,"",AVERAGE(IF({cnd},{post_n})))')
        wc[f"{L(base + 3)}{r}"] = f'=IF({nref}<1,"",{L(base + 2)}{r}-{L(base + 1)}{r})'
    style_row(r, 9, {2: "0", 6: "0"})
cur = cur + 2 + len(SUB) + 1

# K. 일반적 특성
K_TOP = cur
wc.cell(row=cur, column=1, value="K. 일반적 특성 (분석포함=1, 대상자_VO2max 시트) — 원고 표 1용").font = F_SEC
heads = ["변수", "실험 n", "실험 M", "실험 SD", "통제 n", "통제 M", "통제 SD", "표준화 평균차\n(SMD)"]
for j, h in enumerate(heads, start=1):
    hdr_cell(wc, cur + 1, j, h)
wc.row_dimensions[cur + 1].height = 30
for k, (nm, lab) in enumerate((("age", "연령 (세)"), ("height", "신장 (cm)"), ("wt_pre", "사전 체중 (kg)"),
                               ("career", "선수 경력 (년)"))):
    r = cur + 2 + k
    wc.cell(row=r, column=1, value=lab).font = F_BASE
    for gv, base in ((1, 2), (2, 5)):
        cnd = f"(grp={gv})*(inc=1)*ISNUMBER({nm})"
        nref = f"{L(base)}{r}"
        wc[nref] = f"=SUMPRODUCT({cnd})"
        arr(f"{L(base + 1)}{r}", f'=IF({nref}<1,"",AVERAGE(IF({cnd},{nm})))')
        arr(f"{L(base + 2)}{r}", f'=IF({nref}<2,"",STDEV(IF({cnd},{nm})))')
    sp = f"SQRT(((B{r}-1)*D{r}^2+(E{r}-1)*G{r}^2)/(B{r}+E{r}-2))"
    wc[f"H{r}"] = f'=IF(OR(B{r}<2,E{r}<2),"",IF(ROUND({sp},10)=0,"",(C{r}-F{r})/{sp}))'
    style_row(r, 8, {2: "0", 5: "0"})
r = cur + 6
wc.cell(row=r, column=1, value="성별 남 / 여 (명)").font = F_BASE
wc[f"B{r}"] = '=SUMPRODUCT((grp=1)*(inc=1)*(sex=1))&" / "&SUMPRODUCT((grp=1)*(inc=1)*(sex=2))'
wc[f"E{r}"] = '=SUMPRODUCT((grp=2)*(inc=1)*(sex=1))&" / "&SUMPRODUCT((grp=2)*(inc=1)*(sex=2))'
style_row(r, 8, {})
wc.cell(row=r + 1, column=1, value="SMD = 집단 평균 차 ÷ 합동 SD (절댓값 0.1 미만이면 균형이 좋은 편). 기초 균형은 p값보다 SMD로 보고하는 것을 권합니다.").font = F_NOTE
cur = r + 3

# J. 자료 점검
J_TOP = cur
wc.cell(row=cur, column=1, value="J. 자료 점검 요약").font = F_SEC
fr = lambda sh, col: f"{q(sh)}!${col}${FIRST}:${col}${LAST}"
CHECKS = [
    ("대상자_VO2max 입력 인원 (ID 기준)", f'=SUMPRODUCT(--(LEN({fr(S_SUBJ, CS["ID"])})>0))', False),
    ("FMS 시트 입력 ID 수", f'=SUMPRODUCT(--(LEN({fr(S_FMS, CF["ID"])})>0))', False),
    ("FMS_평가자2 입력 ID 수", f'=SUMPRODUCT(--(LEN({fr(S_FMS2, CF2["ID"])})>0))', False),
    ("대상자 시트 경고가 있는 행 수", f'=SUMPRODUCT(--(LEN({fr(S_SUBJ, CS["chk_vo2"])})>0))', True),
    ("FMS 시트 경고가 있는 행 수", f'=SUMPRODUCT(--(LEN({fr(S_FMS, CF["chk_fms"])})>0))', True),
    ("FMS_평가자2 시트 경고가 있는 행 수", f'=SUMPRODUCT(--(LEN({fr(S_FMS2, CF2["chk_fms2"])})>0))', True),
    ("FMS 시트 ID 중 대상자 시트에 없는 ID 수", f'=COUNTIF({fr(S_FMS, CF["group"])},"ID없음")', True),
    ("계산총점 ≠ SPSS총점인 행 수", f'=SUMPRODUCT(--ISNUMBER(SEARCH("≠SPSS총점",{fr(S_FMS, CF["chk_fms"])})))', True),
    ("총점이 정수가 아닌 행 수", f'=SUMPRODUCT(--ISNUMBER(SEARCH("정수 아님",{fr(S_FMS, CF["chk_fms"])})))', True),
    ("하위항목 없이 SPSS총점으로 대체된 FMS 값 수 (사전+사후)",
     f'=COUNTIF({fr(S_FMS, CF["src_pre"])},"SPSS총점")+COUNTIF({fr(S_FMS, CF["src_post"])},"SPSS총점")', True),
    ("clearing test 기록이 3개 미만인 FMS 측정 수 (통증 없음으로 계산됨)",
     f'=COUNTIF({fr(S_FMS, CF["clrn_pre"])},"<3")+COUNTIF({fr(S_FMS, CF["clrn_post"])},"<3")', True),
    ("수정기록 건수", f"=COUNTA({q(S_LOG)}!$C$4:$C$103)", False),
    ("머리글·예시 행이 정렬 등으로 섞였는가 (1이면 Ctrl+Z로 되돌리기; 예시 행을 지워도 1)",
     f'=IF(AND({q(S_SUBJ)}!$A$3="예시",{q(S_FMS)}!$A$3="예시",{q(S_FMS2)}!$A$3="예시",{q(S_ATT)}!$A$3="예시"),0,1)', True),
]
for k, (lab, f, flag) in enumerate(CHECKS):
    r = cur + 1 + k
    wc.cell(row=r, column=1, value=lab).font = F_BASE
    c = wc.cell(row=r, column=2, value=f)
    c.fill, c.alignment, c.number_format = FILL_CALC, CENTER, "0"
    for j in (1, 2):
        wc.cell(row=r, column=j).border = BOX
    wc.cell(row=r, column=1).alignment = WRAP
    if flag:
        wc.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f"B{r}>0"], fill=FILL_WARN))
    wc.row_dimensions[r].height = 28
wc.column_dimensions["A"].width = 34
for j in range(2, 15):
    wc.column_dimensions[L(j)].width = 11

# 흐름 시트 자동 행 연결
ws_flow[f"B{FR['VO2max 분석 인원 (자동: 자동검산 A)']}"] = f"={q(S_CALC)}!B{VE}"
ws_flow[f"C{FR['VO2max 분석 인원 (자동: 자동검산 A)']}"] = f"={q(S_CALC)}!B{VC}"
ws_flow[f"B{FR['FMS 분석 인원 (자동: 자동검산 B)']}"] = f"={q(S_CALC)}!B{FE}"
ws_flow[f"C{FR['FMS 분석 인원 (자동: 자동검산 B)']}"] = f"={q(S_CALC)}!B{FC}"
for direct, auto in (("VO2max 분석 인원 (직접 입력)", "VO2max 분석 인원 (자동: 자동검산 A)"),
                     ("FMS 분석 인원 (직접 입력)", "FMS 분석 인원 (자동: 자동검산 B)")):
    d_, a_ = FR[direct], FR[auto]
    ws_flow[f"F{a_}"] = (f'=IF(COUNT(B{d_}:C{d_})=0,"",IF(AND(B{d_}=B{a_},C{d_}=C{a_}),"같음",'
                         f'"직접 입력값과 자동 집계가 다름: 분석포함·결측 확인"))')
ra, rl = FR["배정 인원"], FR["추적 소실 (사후 측정 못 받음)"]
for lab in ("VO2max 분석 인원 (직접 입력)", "FMS 분석 인원 (직접 입력)"):
    rr = FR[lab]
    ws_flow[f"F{rr}"] = (f'=IF(COUNT(B{ra}:C{ra},B{rr}:C{rr})<4,"",IF(OR(B{rr}>B{ra}-N(B{rl}),C{rr}>C{ra}-N(C{rl})),'
                         f'"분석 인원이 배정−추적소실보다 많음",""))')
re_, ra_ = FR["적격성 평가 인원"], FR["배정 인원"]
ws_flow[f"F{re_}"] = (f'=IF(COUNT(D{re_}:D{ra_})<{ra_ - re_ + 1},"",IF(D{re_}-SUM(D{re_ + 1}:D{ra_ - 1})<>D{ra_},'
                      f'"적격성−배정 전 제외와 배정 합계가 다름",""))')
rs_ = FR["대상자 시트 입력 인원 (자동)"]
ws_flow[f"F{ra_}"] = (f'=IF(COUNT(B{ra_}:C{ra_})<2,"",IF(OR(B{ra_}<>B{rs_},C{ra_}<>C{rs_}),'
                      f'"배정 인원과 대상자 시트 입력 인원이 다름",""))')
ws_flow.conditional_formatting.add(
    "F5:F20", FormulaRule(formula=['ISNUMBER(SEARCH("다름",F5))+ISNUMBER(SEARCH("많음",F5))'], fill=FILL_WARN))
ws_flow.conditional_formatting.add("F5:F20", FormulaRule(formula=['F5="같음"'], fill=FILL_OK))

# ================================================================== 수정기록
ws_log["A1"] = "수정기록 (입력 실수를 고친 경우에만; 원래 값과 근거를 남깁니다)"
ws_log["A1"].font = F_TITLE
LOGH = ["수정일", "시트", "ID", "변수(1행 이름)", "원래 값", "수정한 값", "수정 사유", "근거 문서 (기록지·출력물 등)"]
for j, h in enumerate(LOGH, start=1):
    hdr_cell(ws_log, 2, j, h)
EX_LOG = [dt.date(2026, 10, 20), "대상자_VO2max", "예시", "vo2_post", 25.24, 52.42, "자릿수 전사 실수",
          "가스분석기 출력물 2025-05-30"]
for j, v in enumerate(EX_LOG, start=1):
    c = ws_log.cell(row=3, column=j, value=v)
    c.font, c.fill, c.border = F_EX, FILL_EX, BOX
    c.number_format = "yyyy-mm-dd" if j == 1 else "General"
for r in range(4, 104):
    for j in range(1, 9):
        c = ws_log.cell(row=r, column=j)
        c.fill, c.font, c.border = FILL_IN, F_INPUT, BOX
        if j == 1:
            c.number_format = "yyyy-mm-dd"
for col, w in zip("ABCDEFGH", (11, 16, 8, 14, 10, 10, 30, 32)):
    ws_log.column_dimensions[col].width = w
ws_log.freeze_panes = "A4"

# ================================================================== 안내
g = ws_guide
g["A1"] = "H-IFIT 연구 원자료 입력 양식"
g["A1"].font = Font(name=FONT, size=16, bold=True)
g["A2"] = "12주 고강도 기능성 인터벌 트레이닝 × 대학 태권도 겨루기 선수 (VO2max · FMS)"
g["A2"].font = F_NOTE
GUIDE = [
    ("목적", None),
    ("", "원자료로 포스터 수치를 다시 계산하고(특히 FMS 총점 평균과 t값이 서로 맞지 않는 문제), 통제군 기술통계·자유도·효과크기·95% CI·"
         "대상자 특성을 구해 원고의 표 1·3·4와 연구방법을 완성하기 위한 양식입니다."),
    ("가장 중요한 원칙", None),
    ("!", "원 기록지·측정 출력물의 값을 그대로 입력하세요. 포스터 값과 달라도 원자료를 고치거나 대상자를 빼지 마세요. "
          "'다름'은 그대로 보고할 결과입니다. 입력 실수를 고쳤다면 '수정기록' 시트에 원래 값과 근거를 남깁니다."),
    ("작성 순서", None),
    ("1 연구정보", "배정·IRB·측정 장비·FMS 평가자 등에 답하고 근거(문서·기록 / 기억 / 모름)를 고릅니다."),
    ("2 대상자_VO2max", "1명당 1행(4–33행). ID·집단·분석포함·인적사항·사전/사후 체중과 VO2max."),
    ("3 FMS", "ID를 직접 입력하고 하위항목 원점수를 입력합니다. 집단·분석포함은 ID로 자동 연결되므로 행 순서는 달라도 됩니다."),
    ("3b FMS_평가자2", "(선택) 평가자가 2명이고 두 사람 점수를 따로 갖고 있으면 평가자2 점수를 입력합니다."),
    ("4 출석_훈련기록", "실험군 회기별 출석(1/0), 35행에는 전체 취소된 회기에만 0, 정규훈련 시간·강도(두 집단)."),
    ("5 주차별_프로그램", "주차별 실제 운동 구성(종목·시간·세트·휴식)."),
    ("6 대상자_흐름", "적격성 평가 → 배정 → 중단·추적 소실 → 결과별 분석 인원."),
    ("7 자동검산", "확인만 하면 됩니다. 결과는 기록용이며, 포스터 값과 같은지 다른지 표시됩니다."),
    ("색 구분", None),
    ("노란 칸 + 파란 글씨", "직접 입력하는 칸"),
    ("회색 칸", "자동 계산 칸. 수정하거나 지우지 마세요."),
    ("연회색 기울임 (3행)", "예시 행. 자료 범위 밖이라 분석에서 빠지므로 그대로 두셔도 됩니다."),
    ("주황 칸", "입력 점검 경고. 원 기록을 확인해 입력 실수면 고치고(수정기록에 기재), 기록과 같으면 그대로 둡니다."),
    ("파란 칸 (자동검산 H)", "포스터 값과 다름. 오류 표시가 아니라 기록입니다."),
    ("코딩 규칙", None),
    ("집단", "1 = 실험군(H-IFIT), 2 = 통제군(정규훈련)"),
    ("분석포함", "1 = 포함, 0 = 제외(제외 사유를 목록에서 고르고 상세를 적음). 빈칸은 제외로 처리되며 경고가 뜹니다. 탈락자도 행을 지우지 마세요."),
    ("결과별 포함", "대상자 시트의 분석포함은 VO2max에 적용되고 FMS에도 기본으로 따라갑니다. 결과마다 다르면(예: VO2max만 최대치 판정 "
                  "미충족으로 제외) FMS 시트의 'FMS만 분석포함' 칸에 1/0을 따로 넣습니다."),
    ("성별", "1 = 남, 2 = 여"),
    ("FMS 약어", "DS=Deep Squat, HS=Hurdle Step, ILL=In-line Lunge, SM=Shoulder Mobility, ASLR=Active Straight Leg Raise, "
                 "TSPU=Trunk Stability Push-up, RS=Rotary Stability"),
    ("FMS 원점수", "0–3. 양측 항목(HS, ILL, SM, ASLR, RS)은 좌·우를 따로 입력하면 낮은 점수가 항목 점수가 됩니다. "
                  "동작 중 통증이 있었으면 그 원점수에 0을 입력합니다. 평가자 평균 등으로 소수라면 그대로 입력하세요(경고만 뜸)."),
    ("Clearing test", "1 = 통증 있음 → 해당 항목 0점(SM←어깨 충돌, TSPU←척추 신전, RS←척추 굴곡; SM은 좌·우 중 한쪽이라도 통증이면 1), 0 = 없음. "
                      "검사하지 않았거나 기록이 없으면 빈칸(통증 없음으로 계산하고 자동검산 J에 개수 표시)."),
    ("FMS 총점만 있을 때", "하위항목이 없으면 '총점 기존 SPSS·포스터 값' 칸에만 넣어도 계산됩니다. 이 경우 검증이 아니라 재집계라는 점이 J에 표시됩니다."),
    ("결측값", "측정하지 않았거나 모르는 값은 반드시 빈칸. 0을 넣으면 실제 0으로 계산됩니다."),
    ("날짜", "2025-03-03처럼 하이픈으로 입력(2025.03.03은 날짜로 인식되지 않음)."),
    ("ID", "실명 대신 코드(E01–E13, C01–C12 등). 같은 사람은 모든 시트에서 같은 ID."),
    ("행 관련", "자료 범위는 4–33행(30명)입니다. 행을 삽입·삭제하지 마세요. 정렬은 권하지 않습니다. 꼭 필요하면 4–33행만 선택한 뒤 "
              "[정렬] 대화상자에서 '머리글 포함'을 끄고 정렬하세요(1–3행이 섞이면 대상자가 분석에서 빠짐). FMS·출석 시트는 ID로 연결됩니다."),
    ("포스터 수치가 다르게 나올 수 있는 이유 (중립적 후보)", None),
    ("", "평가자 점수의 평균 사용, 좌우 점수를 평균·합산한 다른 채점 규칙, 실제 분석 인원 차이(탈락자), 다른 버전의 자료 파일, "
         "공변량을 넣은 분석(ANCOVA), 포스터로 옮겨 적을 때의 전사 오류. 어느 것인지는 자료와 연구정보 답으로 확인합니다."),
    ("Excel·SPSS", None),
    ("Excel", "파일을 열면 자동으로 계산됩니다. 위쪽에 '편집 사용' 버튼이 보이면 눌러야 계산 결과가 나타납니다."),
    ("SPSS", "1행이 SPSS 변수명입니다. 복사본에서 2·3행을 지우고, 불러올 때 범위를 마지막 대상자 행까지만 지정하세요"
             "(빈 행의 \"\"가 문자 변수로 읽힐 수 있음). 불러온 뒤 숫자 열이 문자형이면 ALTER TYPE d_vo2 pct_vo2 d_wt (F8.3). 처럼 바꿉니다. "
             "group_lbl, chk_vo2, chk_fms 같은 문자 열은 분석에서 빼세요."),
    ("보낼 것", None),
    ("", "작성한 이 파일 + 있으면 기존 SPSS 파일(.sav)과 출력(.spv) + IRB 승인번호. 받으면 최종 분석(비모수 검정, ICC 등)과 원고 수정을 진행합니다."),
]
r = 4
for a, b in GUIDE:
    if b is None:
        c = g.cell(row=r, column=1, value=a)
        c.font, c.fill = F_H2, FILL_HDR
        g.cell(row=r, column=2).fill = FILL_HDR
        g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    else:
        c1 = g.cell(row=r, column=1, value=a)
        c1.font, c1.alignment = F_BOLD, WRAP
        c2 = g.cell(row=r, column=2, value=b)
        c2.font, c2.alignment = (F_RED if a == "!" else F_BASE), WRAP
        fills = {"노란 칸 + 파란 글씨": FILL_IN, "회색 칸": FILL_CALC, "연회색 기울임 (3행)": FILL_EX,
                 "주황 칸": FILL_WARN, "파란 칸 (자동검산 H)": FILL_DIFF}
        if a in fills:
            c1.fill = fills[a]
        if a == "노란 칸 + 파란 글씨":
            c1.font = Font(name=FONT, size=10, bold=True, color="0000FF")
        g.row_dimensions[r].height = 16 if len(b) < 70 else (32 if len(b) < 150 else 46)
    r += 1
g.column_dimensions["A"].width = 24
g.column_dimensions["B"].width = 115

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90

# ================================================================== 검산용 가상자료 (경계 사례 포함)
if SAMPLE:
    rnd = random.Random(20261016)
    ids = [f"E{k:02d}" for k in range(1, 14)] + [f"C{k:02d}" for k in range(1, 13)]
    for k, pid in enumerate(ids):
        row = FIRST + k
        grp = 1 if pid.startswith("E") else 2

        def put(n, v, row=row):
            ws_subj[f"{CS[n]}{row}"] = v
        put("ID", pid); put("group", grp)
        put("include", 0 if pid == "C12" else 1)
        if pid == "C12":
            put("excl_reason", "사후 미측정")
        if pid == "C11":
            put("include", None)                               # 분석포함 빈칸 → 제외 + 경고
        put("sex", 1 if k % 3 else 2); put("age", 19 + k % 4)
        put("height", round(rnd.uniform(165, 190), 1)); put("career", round(rnd.uniform(5, 12), 1))
        wpre = round(rnd.uniform(55, 85), 1)
        put("wt_pre", wpre)
        vp = round(rnd.gauss(48.3, 4), 2)
        put("vo2_pre", vp)
        put("vo2abs_pre", round(vp * wpre))
        put("rer_pre", 1.12); put("hrmax_pre", 196); put("rpe_pre", 18); put("crit_pre", 1)
        put("date_pre", dt.date(2025, 3, 3))
        if pid not in ("E13", "C12"):
            put("wt_post", round(wpre + rnd.uniform(-2, 1), 1))
            put("vo2_post", round(vp + (rnd.gauss(4, 2.3) if grp == 1 else rnd.gauss(0.5, 1.6)), 2))
            put("date_post", dt.date(2025, 5, 30))
        if pid == "E12":
            put("wt_post", None)                               # 사후 체중 빈칸
            put("vo2abs_post", 3500)                           # 절대값만 있음 → 오류 없어야 함
        if pid == "E11":
            put("rer_pre", 1.02); put("hrmax_pre", 170); put("rpe_pre", 15)   # 판정충족=1인데 기준 미달
            put("include", 0); put("excl_reason", "최대치 판정 미충족")          # VO2만 제외
        if pid == "E10":
            put("date_pre", "2025.03.03")                      # 텍스트 날짜
    fms_rows = list(reversed(ids)) + ["X99"]                   # 역순 입력(ID 연결 검증), X99=없는 ID
    for k, pid in enumerate(fms_rows):
        row = FIRST + k
        grp = 1 if pid.startswith("E") else 2
        ws_fms[f"A{row}"] = pid
        if pid == "X99":
            continue
        for t in ("pre", "post"):
            if pid == "C12" and t == "post":
                continue
            for n in move_names(t):
                base = rnd.choice([1, 2, 2, 2, 3])
                if t == "post" and grp == 1 and base < 3 and rnd.random() < 0.4:
                    base += 1
                ws_fms[f"{CF[n]}{row}"] = base
            for n in clr_names(t):
                if pid == "E03":
                    continue                                    # clearing 미기록 → 원점수로 계산
                ws_fms[f"{CF[n]}{row}"] = 1 if rnd.random() < 0.05 else 0
        if pid == "E05":
            ws_fms[f"{CF['fms_spss_pre']}{row}"] = 14.5
        if pid == "E11":
            ws_fms[f"{CF['include_fms']}{row}"] = 1                      # FMS는 분석에 남김
        if pid == "C02":                                        # 하위항목 없이 SPSS총점만
            for t in ("pre", "post"):
                for n in move_names(t) + clr_names(t):
                    ws_fms[f"{CF[n]}{row}"] = None
            ws_fms[f"{CF['fms_spss_pre']}{row}"] = 13
            ws_fms[f"{CF['fms_spss_post']}{row}"] = 14
    for k, pid in enumerate(["E01", "E02", "E03", "E04", "C01", "C03", "C04"]):
        row = FIRST + k
        ws_fms2[f"A{row}"] = pid
        src_row = FIRST + fms_rows.index(pid)
        for t in ("pre", "post"):
            for n in move_names(t) + clr_names(t):
                v = ws_fms[f"{CF[n]}{src_row}"].value
                if v is not None and n in move_names(t) and rnd.random() < 0.15:
                    v = max(0, min(3, v + rnd.choice([-1, 1])))
                ws_fms2[f"{CF2[n]}{row}"] = v
    for k, pid in enumerate(ids[:13]):
        row = FIRST + k
        ws_att[f"A{row}"] = pid
        for s in range(1, NSESS + 1):
            ws_att[f"{CA[f's{s:02d}']}{row}"] = 0 if rnd.random() < 0.06 else 1
    for s in range(1, NSESS + 1):
        ws_att[f"{CA[f's{s:02d}']}{HELD_ROW}"] = 0 if s == 20 else 1

import re as _re
from openpyxl.utils import column_index_from_string as _ci
for _n in wb.defined_names:
    _m = _re.fullmatch(r"([A-Za-z]{1,3})(\d+)", _n)
    assert not (_m and _ci(_m.group(1).upper()) <= 16384), f"이름이 셀 주소와 같음: {_n}"
    assert not _re.fullmatch(r"[RrCc]\d*([RrCc]\d*)?", _n), f"R1C1 형태 이름: {_n}"
wb.save(OUT)
print("saved", OUT, "| rows: VE", VE, "FE", FE, "AE", AE, "VI", VI, "FI", FI, "AI", AI, "VB", VB, "FB", FB,
      "H", H_TOP, "K", K_TOP, "J", J_TOP)
