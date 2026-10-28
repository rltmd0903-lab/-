"""표 1(대상자의 일반적 특성과 사전 종속변인) 계산 시트.

사용: python3 build_t1.py <출력.xlsx> [--sample]
"""
import sys
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

OUT = sys.argv[1]
SAMPLE = "--sample" in sys.argv
FONT = "맑은 고딕"
F = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)
IN = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="E7E6E6")
HDR = PatternFill("solid", fgColor="1F4E78")
EX = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
C = Alignment(horizontal="center", vertical="center", wrap_text=True)

FIRST, LAST = 5, 34
wb = Workbook()
wi = wb.active
wi.title = "입력"
wo = wb.create_sheet("표1")

wi["A1"] = ("노란 칸에 대상자별 값을 입력하세요. 성별: 1=남, 2=여. 모르는 값은 빈칸. 사전 VO2max와 FMS 총점은 사전사후 입력 시트의 "
            "사전값과 같은 값을 넣으면 됩니다. 결과는 '표1' 시트에 원고 표 1 형식으로 계산됩니다.")
wi["A1"].font = F(size=9, color="595959")
wi["A1"].alignment = Alignment(wrap_text=True, vertical="center")
wi.merge_cells("A1:I1")
wi.row_dimensions[1].height = 32
COLS = [("ID", 8), ("집단\n1=실험\n2=통제", 8), ("성별\n1=남\n2=여", 8), ("연령\n(세)", 8), ("신장\n(cm)", 8),
        ("체중\n(kg)", 8), ("선수 경력\n(년)", 9), ("사전 VO2max\n(mL/kg/min)", 11), ("사전 FMS\n총점 (점)", 9)]
for j, (h, w) in enumerate(COLS, start=1):
    for r in (2, 3):
        c = wi.cell(row=r, column=j)
        c.fill, c.border = HDR, BOX
    wi.merge_cells(start_row=2, start_column=j, end_row=3, end_column=j)
    c = wi.cell(row=2, column=j, value=h)
    c.font, c.alignment = F(bold=True, color="FFFFFF", size=9), C
    wi.column_dimensions[c.column_letter].width = w
wi.row_dimensions[2].height = 26
ex = ["예시", "(예시)", 1, 21, 175.2, 68.4, 9.5, 47.85, 14]
ids = [f"E{k:02d}" for k in range(1, 14)] + [f"C{k:02d}" for k in range(1, 13)]
for r in [4] + list(range(FIRST, LAST + 1)):
    for j in range(1, 10):
        c = wi.cell(row=r, column=j)
        c.border, c.alignment = BOX, Alignment(horizontal="center", vertical="center")
        c.fill = EX if r == 4 else IN
        c.font = F(italic=True, color="7F7F7F") if r == 4 else F(color="0000FF")
        c.number_format = {4: "0", 5: "0.0", 6: "0.0", 7: "0.0", 8: "0.00", 9: "0"}.get(j, "General")
for j, v in enumerate(ex, start=1):
    wi.cell(row=4, column=j, value=v)
for k, pid in enumerate(ids):
    wi.cell(row=FIRST + k, column=1, value=pid)
    wi.cell(row=FIRST + k, column=2, value=1 if pid.startswith("E") else 2)
dv = DataValidation(type="list", formula1='"1,2"', allow_blank=True, showErrorMessage=True, errorTitle="코드", error="1 또는 2")
dv.add(f"B{FIRST}:C{LAST}")
wi.add_data_validation(dv)
wi.freeze_panes = f"C{FIRST}"

# ------------------------------------------------------------------ 표1
G = f"'입력'!$B${FIRST}:$B${LAST}"
rng = lambda col: f"'입력'!${col}${FIRST}:${col}${LAST}"
wo["A1"] = "표 1. 대상자의 일반적 특성과 사전 종속변인 (원고에 그대로 옮길 값은 '논문 표기' 열)"
wo["A1"].font = F(bold=True, size=12)
wo["A2"] = "값은 평균±표준편차. p는 독립표본 t검정(등분산 가정), 성별은 Fisher의 정확검정(양측). 값이 비어 있는 대상자는 그 변인에서만 제외됩니다."
wo["A2"].font = F(size=9, color="595959")
heads = ["변인", "실험군 n", "실험군 M", "실험군 SD", "통제군 n", "통제군 M", "통제군 SD", "t", "df", "p",
         "실험군\n(논문 표기)", "통제군\n(논문 표기)", "p\n(논문 표기)"]
widths = [22, 8, 9, 9, 8, 9, 9, 8, 6, 8, 13, 13, 10]
for j, (h, w) in enumerate(zip(heads, widths), start=1):
    c = wo.cell(row=4, column=j, value=h)
    c.font, c.fill, c.alignment, c.border = F(bold=True, color="FFFFFF", size=9), HDR, C, BOX
    wo.column_dimensions[c.column_letter].width = w
wo.row_dimensions[4].height = 30
VARS = [("연령 (세)", "D", "0.0"), ("신장 (cm)", "E", "0.0"), ("체중 (kg)", "F", "0.0"), ("선수 경력 (년)", "G", "0.0"),
        ("VO2max (mL/kg/min)", "H", "0.00"), ("FMS 총점 (점)", "I", "0.00")]
r = 5
for name, col, fmt in VARS:
    x = rng(col)
    wo[f"A{r}"] = name
    for g, base in ((1, 2), (2, 5)):
        cnd = f"({G}={g})*ISNUMBER({x})"
        n, m, sd = (chr(64 + base), chr(65 + base), chr(66 + base))
        wo[f"{n}{r}"] = f"=SUMPRODUCT({cnd})"
        wo[f"{m}{r}"] = ArrayFormula(f"{m}{r}", f'=IF({n}{r}<1,"",AVERAGE(IF({cnd},{x})))')
        wo[f"{sd}{r}"] = ArrayFormula(f"{sd}{r}", f'=IF({n}{r}<2,"",STDEV(IF({cnd},{x})))')
    sp = f"SQRT(((B{r}-1)*D{r}^2+(E{r}-1)*G{r}^2)/(B{r}+E{r}-2))"
    wo[f"H{r}"] = f'=IF(OR(B{r}<2,E{r}<2),"",IF(ROUND({sp},10)=0,"",(C{r}-F{r})/({sp}*SQRT(1/B{r}+1/E{r}))))'
    wo[f"I{r}"] = f'=IF(H{r}="","",B{r}+E{r}-2)'
    wo[f"J{r}"] = f'=IF(H{r}="","",TDIST(ABS(H{r}),I{r},2))'
    dec = "0.00" if fmt == "0.00" else "0.0"
    wo[f"K{r}"] = f'=IF(C{r}="","",TEXT(C{r},"{dec}")&"±"&IF(D{r}="","",TEXT(D{r},"{dec}")))'
    wo[f"L{r}"] = f'=IF(F{r}="","",TEXT(F{r},"{dec}")&"±"&IF(G{r}="","",TEXT(G{r},"{dec}")))'
    wo[f"M{r}"] = f'=IF(J{r}="","",IF(J{r}<0.001,"<.001",TEXT(J{r},".000")))'
    for j in range(1, 14):
        c = wo.cell(row=r, column=j)
        c.border, c.alignment = BOX, Alignment(horizontal="center", vertical="center")
        if j > 1:
            c.fill, c.font = CALC, F(size=9)
            c.number_format = {2: "0", 5: "0", 8: "0.000", 9: "0", 10: "0.0000"}.get(j, fmt if j in (3, 4, 6, 7) else "@")
    r += 1
# 성별: Fisher 정확검정
SX = rng("C")
wo[f"A{r}"] = "성별 (남/여, 명)"
wo[f"B{r}"] = f"=SUMPRODUCT(({G}=1)*({SX}=1))"     # a: 실험 남
wo[f"C{r}"] = f"=SUMPRODUCT(({G}=1)*({SX}=2))"     # b: 실험 여
wo[f"E{r}"] = f"=SUMPRODUCT(({G}=2)*({SX}=1))"     # c: 통제 남
wo[f"F{r}"] = f"=SUMPRODUCT(({G}=2)*({SX}=2))"     # d: 통제 여
a, b, c_, d = f"B{r}", f"C{r}", f"E{r}", f"F{r}"
N, r1, c1 = f"({a}+{b}+{c_}+{d})", f"({a}+{b})", f"({a}+{c_})"
hk = f"IFERROR(HYPGEOMDIST(ROW($1:$40)-1,{r1},{c1},{N}),0)"
wo[f"J{r}"] = ArrayFormula(f"J{r}", f'=IF(OR({N}=0,{c1}=0,{c1}={N}),"",MIN(1,SUM(IF({hk}<=HYPGEOMDIST({a},{r1},{c1},{N})*(1+1E-7),{hk},0))))')
wo[f"K{r}"] = f'={a}&" / "&{b}'
wo[f"L{r}"] = f'={c_}&" / "&{d}'
wo[f"M{r}"] = f'=IF(J{r}="","—",IF(J{r}<0.001,"<.001",TEXT(J{r},".000")))'
for j in range(1, 14):
    cc = wo.cell(row=r, column=j)
    cc.border, cc.alignment = BOX, Alignment(horizontal="center", vertical="center")
    if j > 1:
        cc.fill, cc.font = CALC, F(size=9)
        cc.number_format = "0.0000" if j == 10 else "General"
wo.cell(row=r + 2, column=1, value=("성별이 모두 같으면(예: 전원 남자) p는 '—'로 표시됩니다. 표 1에서 성별 행 대신 '전원 남자'라고 주석에 적어도 됩니다. "
                                    "사전 VO2max나 FMS의 p < .05이면 원고의 공분산분석 결과를 함께 보고하세요.")).font = F(size=9, color="595959")
wo.freeze_panes = "B5"

if SAMPLE:
    rnd = random.Random(3)
    for k in range(25):
        rr = FIRST + k
        wi.cell(row=rr, column=3, value=1 if rnd.random() < 0.7 else 2)
        wi.cell(row=rr, column=4, value=rnd.randint(19, 24))
        wi.cell(row=rr, column=5, value=round(rnd.gauss(175, 6), 1))
        wi.cell(row=rr, column=6, value=round(rnd.gauss(68, 7), 1))
        wi.cell(row=rr, column=7, value=round(rnd.uniform(5, 12), 1))
        if k != 4:
            wi.cell(row=rr, column=8, value=round(rnd.gauss(48, 3), 2))
        wi.cell(row=rr, column=9, value=rnd.randint(11, 17))

wb.save(OUT)
print("saved", OUT)
