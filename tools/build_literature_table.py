"""선행연구 VO2max·FMS 사전/사후 값 정리 파일 생성.

사용: python3 build_lit.py <출력.xlsx>
수치는 초록·2차 요약에서 확인한 값이며, '원문 확인'은 원문 표에서만 확인 가능한 값이다.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
FONT = "맑은 고딕"
F = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)
HDR = PatternFill("solid", fgColor="1F4E78")
CALC = PatternFill("solid", fgColor="E7E6E6")
CHECK = PatternFill("solid", fgColor="FFF2CC")   # 원문 확인 필요
EXCL = PatternFill("solid", fgColor="F2F2F2")    # 프로토콜 불일치(비교 제외)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
NA = "원문 확인"

URL = {
    "batra": "https://www.researchgate.net/publication/311446299_Effect_of_high_intensity_interval_training_on_cardiopulmonary_function_in_Taekwon-do_ITF_athletes",
    "xu": "https://onlinelibrary.wiley.com/doi/10.1002/ejsc.70050",
    "wr2019": "https://pubmed.ncbi.nlm.nih.gov/31288289/",
    "monks": "https://www.minervamedica.it/en/journals/sports-med-physical-fitness/article.php?cod=R40Y2017N10A1252",
    "song": "https://pubmed.ncbi.nlm.nih.gov/39228785/",
    "indo": "https://digilib.unimed.ac.id/68226/3/6221210001_Abstrak.pdf",
    "norway": "https://www.frontiersin.org/journals/physiology/articles/10.3389/fphys.2025.1661237/full",
    "bulgaria": "https://zenodo.org/records/14729990",
    "iran": "https://brieflands.com/articles/asjsm-94796",
    "taper": "https://www.iat.uni-leipzig.de/datenbanken/iks/dfb-fechten/Record/4053069",
    "vasc": "https://www.ovid.com/jnls/nsca-jscr/fulltext/10.1519/jsc.0000000000003255~effects-of-high-intensity-interval-training-in-combat-sports",
    "ma26": "https://www.frontiersin.org/journals/nutrition/articles/10.3389/fnut.2026.1792680/full",
    "ma23": "https://www.ncbi.nlm.nih.gov/pmc/articles/PMC10279791/",
    "bodden": "https://pubmed.ncbi.nlm.nih.gov/23860293/",
    "kiesel11": "https://www.functionalmovement.com/files/articles/79a_fms_intervention_10.07.pdf",
    "peerj": "https://peerj.com/articles/16878.pdf",
    "skate": "https://lida.sport-iat.de/ta/Record/4092861?lng=en",
    "fire": "https://www.researchgate.net/publication/314080729_Active_duty_firefighters_can_improve_Functional_Movement_Screen_FMS_scores_following_an_8-week_individualized_client_workout_program",
    "frost": "https://www.backfitpro.com/medical-scientific-articles/2015/%5B26%5D-Frost,D.M.-(2015)-FMS-scores-change-with-performer's-knowledge-of-the-grading-criteria-%5BJ.-Strength-Cond.-Res.%5D.pdf",
    "volley": "https://journalofsportsmedicine.org/abstract/482/eng",
    "kim26": "https://www.kci.go.kr/kciportal/landing/article.kci?arti_id=ART003346567",
    "pnf": "https://www.kci.go.kr/kciportal/ci/sereArticleSearch/ciSereArtiView.kci?sereArticleSearchBean.artiId=ART002194434",
    "pollen": "https://researchdiscovery.drexel.edu/esploro/outputs/journalArticle/Do-Normative-Composite-Scores-on-the/991019167434504721",
    "schneiders": "https://www.researchgate.net/publication/51452094_Functional_Movement_Screen_Normative_values_in_a_young_active_population",
    "gaelic": "https://www.sciencedirect.com/science/article/abs/pii/S1466853X13001089",
    "adol": "https://pmc.ncbi.nlm.nih.gov/articles/PMC4325285/",
    "college": "https://ijspt.scholasticahq.com/article/19443-functional-movement-screen-detected-asymmetry-normative-values-among-college-aged-students",
    "hs": "https://www.researchgate.net/publication/278330710_Functional_movement_screen_normative_values_and_validity_in_high_school_athletes_can_the_FMS_be_used_as_a_predictor_of_injury",
    "maleki": "https://www.nature.com/articles/s41598-025-12371-2",
    "highrisk": "https://ijspt.scholasticahq.com/article/31001-effective-interventions-for-improving-functional-movement-screen-scores-among-high-risk-athletes-a-systematic-review",
    "teyhen": "https://www.jospt.org/doi/10.2519/jospt.2012.3838",
    "leejin": "https://koreascience.kr/article/JAKO201734549518480.pdf",
}

wb = Workbook()


def header(ws, row, labels, widths):
    for j, (h, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.alignment, c.border = F(bold=True, color="FFFFFF", size=9), HDR, CEN, BOX
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[row].height = 32


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = F(bold=True, size=13)
    ws["A2"] = note
    ws["A2"].font = F(size=9, color="595959")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[2].height = 30


# 사전·사후 표 공통 열
PP_HEAD = ["연구", "대상", "n", "중재·기간", "측정 프로토콜", "집단", "사전 M", "사전 SD", "사후 M", "사후 SD",
           "변화량\n(자동)", "변화율\n(자동)", "보고된 변화·결과", "근거 수준 / 비고", "출처"]
PP_W = [22, 20, 6, 24, 18, 10, 8, 8, 8, 8, 9, 9, 26, 26, 10]


def pp_sheet(ws, rows):
    header(ws, 4, PP_HEAD, PP_W)
    r = 5
    for row in rows:
        (study, subj, n, interv, proto, grp, pm, psd, qm, qsd, reported, note, key) = row[:13]
        flag = row[13] if len(row) > 13 else None
        vals = [study, subj, n, interv, proto, grp, pm, psd, qm, qsd]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = F(size=9), BOX
            c.alignment = CEN if j in (3, 6, 7, 8, 9, 10) else WRAP
            if v == NA:
                c.fill = CHECK
            if j in (7, 8, 9, 10) and isinstance(v, (int, float)):
                c.number_format = "0.00"
        ws.cell(row=r, column=11, value=f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(I{r})),I{r}-G{r},"")')
        ws.cell(row=r, column=12, value=f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(I{r})),IF(G{r}<>0,(I{r}-G{r})/G{r},""),"")')
        for j, fmt in ((11, "+0.00;-0.00;0.00"), (12, "+0.0%;-0.0%;0.0%")):
            c = ws.cell(row=r, column=j)
            c.fill, c.font, c.border, c.alignment, c.number_format = CALC, F(size=9), BOX, CEN, fmt
        for j, v in ((13, reported), (14, note)):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border, c.alignment = F(size=9), BOX, WRAP
        c = ws.cell(row=r, column=15, value="링크")
        c.hyperlink, c.font, c.border, c.alignment = URL[key], F(size=9, color="0563C1", underline="single"), BOX, CEN
        if flag == "excl":
            for j in range(1, 16):
                if ws.cell(row=r, column=j).fill != CHECK:
                    ws.cell(row=r, column=j).fill = EXCL
        ws.row_dimensions[r].height = 44
        r += 1
    ws.freeze_panes = "B5"
    return r


# ------------------------------------------------------------------ 0. 안내
ws = wb.active
ws.title = "안내"
ws["A1"] = "선행연구 VO2max·FMS 사전/사후 값 정리"
ws["A1"].font = F(bold=True, size=14)
lines = [
    ("용도", "재분석 결과가 선행연구 범위 안에 있는지 확인하고, 논의에서 선행연구와 비교하기 위한 참고 자료입니다. "
             "이 연구의 결과(원고 표 3·4)에 넣는 값이 아닙니다. 이 연구의 결과는 '사전사후_입력시트'에 입력한 원자료에서 계산합니다."),
    ("수치 출처", "논문 초록과 2차 요약(메타분석 표, 서지 데이터베이스)에서 확인한 값입니다. 원문 사이트에 직접 접속하지 못했으므로 "
                 "논문에 인용하기 전에 원문 표로 다시 확인하세요."),
    ("노란 칸", "'원문 확인' = 초록에 수치가 없어 원문 표에서만 확인할 수 있는 값"),
    ("회색 행", "측정 프로토콜이 이 연구와 달라 직접 비교에서 제외하는 연구(예: 셔틀런 현장검사로 추정한 VO2max)"),
    ("회색 칸(자동)", "변화량(사후−사전)과 변화율은 사전·사후 평균이 모두 있을 때 자동 계산됩니다"),
    ("이 연구의 프로토콜", "VO2max: 트레드밀 점증부하검사(호흡가스 직접측정). FMS: Cook 등(2014) 표준 7항목, 0–3점, 좌우 중 낮은 점수, "
                        "clearing test 통증 시 0점, 총점 0–21점. 이 기준과 같은 방식으로 측정한 연구를 우선 비교 대상으로 했습니다."),
    ("시트 구성", "VO2max_중재 · VO2max_기준값 · FMS총점_중재 · FMS총점_기준값 · FMS하위항목 · 메타분석 · 참고범위_요약"),
]
for i, (a, b) in enumerate(lines, start=3):
    ws.cell(row=i, column=1, value=a).font = F(bold=True)
    c = ws.cell(row=i, column=2, value=b)
    c.font, c.alignment = F(), WRAP
    ws.cell(row=i, column=1).alignment = WRAP
    ws.row_dimensions[i].height = 46 if len(b) > 80 else 22
ws.cell(row=5, column=1).fill = CHECK
ws.cell(row=6, column=1).fill = EXCL
ws.cell(row=7, column=1).fill = CALC
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 110

# ------------------------------------------------------------------ 1. VO2max 중재
ws = wb.create_sheet("VO2max_중재")
title(ws, "VO2max 중재연구 사전/사후 (mL/kg/min)", "같은 행의 사전·사후는 같은 집단의 평균±표준편차입니다. 변화율만 보고된 연구는 '보고된 변화·결과' 열에 적었습니다.")
VO2_ROWS = [
    ("Batra & Zatoń (2016), J Combat Sports Martial Arts", "ITF 태권도 수련자(남)", 20, "8주, 주 2회 인터벌(돌려차기 30초/휴식 90초) 추가",
     "트레드밀 점증부하, 직접측정", "실험군", 50.13, 3.81, 53.16, NA, "유의하게 증가(p ≤ .05)", "이 연구와 프로토콜 가장 유사. 사후 SD 원문 확인", "batra"),
    ("Batra & Zatoń (2016)", "ITF 태권도 수련자(남)", NA, "정규 수련", "트레드밀 점증부하, 직접측정", "통제군", NA, NA, NA, NA,
     "유의한 변화 없음", "통제군 수치 원문 확인", "batra"),
    ("Xu 외 (2025), Eur J Sport Sci", "청소년 태권도 16–20세", 30, "HIIT 30초:10초", "VO2peak", "HIIT-30S", NA, NA, NA, NA,
     "+4.72%", "절대값 원문 확인", "xu"),
    ("Xu 외 (2025)", "청소년 태권도 16–20세", NA, "HIIT 20초:10초", "VO2peak", "HIIT-20S", NA, NA, NA, NA, "+1.96%", "절대값 원문 확인", "xu"),
    ("Xu 외 (2025)", "청소년 태권도 16–20세", NA, "정규훈련", "VO2peak", "통제군", NA, NA, NA, NA, "+0.47%", "절대값 원문 확인", "xu"),
    ("운동:휴식 비율별 HIIT (2019)", "남자 태권도 15–18세", 47, "HIIT 3개 비율 + 통제군", "VO2max", "실험군", NA, NA, NA, NA,
     "+4.84%. 1:4 비율 집단만 통제군 대비 유의", "집단별 절대값 원문 확인", "wr2019"),
    ("Monks 외 (2017), J Sports Med Phys Fitness", "대학 태권도 선수(남녀)", 33, "4주 11회, HIIT(16명) vs 고강도 지속달리기(17명)",
     "VO2max", "HIIT", NA, NA, NA, NA, "두 집단 모두 향상(초록 판본 간 서술 차이 있음)", "수치 원문 확인", "monks"),
    ("Song & Sheykhlouvand (2024), J Sports Sci Med", "태권도 선수", NA, "약 8주 16회, 기술 특이 인터벌 vs 짧은 질주 인터벌 vs 통제",
     "VO2max, 심박출량 등", "HIIT", NA, NA, NA, NA, "두 HIIT 모두 향상, 기술 특이 인터벌이 더 큼", "수치·표본 수 원문 확인", "song"),
    ("인도네시아 겨루기 HIIT", "겨루기 선수", 6, "HIIT, 단일군", "셔틀런(MFT) 추정", "실험군", 27.26, NA, 30.43, NA,
     "유의하게 증가", "현장검사 추정치라 직접 비교 제외", "indo", "excl"),
]
end = pp_sheet(ws, VO2_ROWS)

# ------------------------------------------------------------------ 2. VO2max 기준값
ws = wb.create_sheet("VO2max_기준값")
title(ws, "태권도 선수 VO2max 기준값 (사전값 비교용, mL/kg/min)", "중재 없이 측정한 횡단 자료입니다.")
BASE_HEAD = ["연구", "대상", "n", "측정 프로토콜", "구분", "평균", "SD", "비고", "출처"]
header(ws, 4, BASE_HEAD, [26, 24, 6, 18, 10, 9, 9, 30, 10])
VO2_BASE = [
    ("노르웨이 ITF 국가대표 (2025), Front Physiol", "국가대표", NA, "실험실 직접측정", "남", 58.8, 7.3, "", "norway"),
    ("노르웨이 ITF 국가대표 (2025)", "국가대표", NA, "실험실 직접측정", "여", 50.5, 4.1, "", "norway"),
    ("불가리아 국가대표 (2021)", "국가대표", 8, "트레드밀", "남", 58.2, NA, "", "bulgaria"),
    ("불가리아 국가대표 (2021)", "국가대표", 6, "트레드밀", "여", 46.0, NA, "", "bulgaria"),
    ("이란 엘리트 여자 선수", "엘리트", 33, NA, "여", 48.95, 7.11, "", "iran"),
    ("테이퍼링 연구", "태권도 선수", NA, NA, "남", 53.5, NA, "약값", "taper"),
    ("테이퍼링 연구", "태권도 선수", NA, NA, "여", 45.5, NA, "약값", "taper"),
]


def plain_rows(ws, rows, start=5, num_cols=(6, 7)):
    r = start
    for row in rows:
        *vals, key = row
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.border = F(size=9), BOX
            c.alignment = CEN if j >= 3 and j != len(vals) else WRAP
            if v == NA:
                c.fill = CHECK
            if j in num_cols and isinstance(v, (int, float)):
                c.number_format = "0.00"
        c = ws.cell(row=r, column=len(vals) + 1, value="링크")
        c.hyperlink, c.font, c.border, c.alignment = URL[key], F(size=9, color="0563C1", underline="single"), BOX, CEN
        ws.row_dimensions[r].height = 30
        r += 1
    ws.freeze_panes = f"B{start}"
    return r


plain_rows(ws, VO2_BASE)

# ------------------------------------------------------------------ 3. FMS 총점 중재
ws = wb.create_sheet("FMS총점_중재")
title(ws, "FMS 총점 중재연구 사전/사후 (0–21점)", "FMS 총점은 Cook 등의 표준 7항목 합계입니다.")
FMS_ROWS = [
    ("Bodden 외 (2015), J Strength Cond Res", "MMA 선수(준프로)", 25, "8주, 교정운동 주 4회 + 평소 훈련", "FMS 표준", "실험군",
     13.25, 0.87, NA, NA, "0→8주 유의하게 증가, 대부분 4주 안에", "사전값은 2차 자료 기준", "bodden"),
    ("Bodden 외 (2015)", "MMA 선수(준프로)", NA, "평소 훈련", "FMS 표준", "통제군", 13.23, 0.80, NA, NA, "유의한 증가 없음", "사후값 원문 확인", "bodden"),
    ("Kiesel 외 (2011), Scand J Med Sci Sports", "NFL 라인맨", NA, "7주 비시즌 프로그램", "FMS 표준", "라인맨", 11.8, NA, 14.8, NA,
     "유의하게 증가", "SD 원문 확인. 통제군도 증가했다는 2차 보고 있음", "kiesel11"),
    ("Kiesel 외 (2011)", "NFL 비라인맨", NA, "7주 비시즌 프로그램", "FMS 표준", "비라인맨", 13.3, NA, 16.3, NA,
     "유의하게 증가", "SD 원문 확인", "kiesel11"),
    ("기능 교정 훈련 (PeerJ)", "대학생", 15, "기능 교정 훈련", "FMS 표준", "실험군", 15.73, 1.58, 18.07, 1.79, "유의하게 증가(p = .001)",
     "비선수. 대조군(13명) 수치 원문 확인", "peerj"),
    ("빙상 선수 8주 FMS 훈련 (JHSE 2025)", "청소년 스피드스케이팅", NA, "8주 FMS 기반 훈련", "FMS 표준", "실험군", 15.38, 0.92, 19.13, 0.64,
     "유의하게 증가", "", "skate"),
    ("소방관 8주 개별 운동", "현역 소방관", NA, "8주 개별 프로그램, 단일군", "FMS 표준", "단일군", 11.3, 2.2, 14.7, 2.8, "p = .013",
     "비선수, 통제군 없음", "fire"),
    ("Frost 외 (2015), J Strength Cond Res", "소방관", NA, "훈련이 아니라 채점 기준을 알려준 뒤 재검사", "FMS 표준", "단일군", 14.1, NA, 16.7, NA,
     "증가", "훈련 효과가 아닌 학습효과 연구. 학습효과 크기 참고용", "frost"),
    ("엘리트 여자 배구 교정운동", "엘리트 여자 배구", NA, "12주 교정운동 주 3회", "FMS 표준", "실험군", NA, NA, NA, NA, "+19.6%",
     "절대값 원문 확인", "volley"),
    ("김현륜·이상일 (2026), 대한무도학회지", "태권도 전공자", NA, "12주 HIIT", "FMS 표준", "실험군", NA, NA, NA, NA,
     "집단 간 유의한 차이. DS·HS·ILL 유의", "수치 원문 확인", "kim26"),
    ("PNF 복합운동 (KCI)", "중·고교 남자 태권도 선수", NA, "PNF 복합운동", "FMS 표준", "실험군", NA, NA, NA, NA,
     "총점과 SM 제외 6개 항목 유의하게 향상", "수치 원문 확인", "pnf"),
]
pp_sheet(ws, FMS_ROWS)

# ------------------------------------------------------------------ 4. FMS 총점 기준값
ws = wb.create_sheet("FMS총점_기준값")
title(ws, "FMS 총점 기준값 (사전값 비교용, 0–21점)", "중재 없이 측정한 횡단 자료입니다.")
header(ws, 4, BASE_HEAD, [26, 24, 6, 18, 10, 9, 9, 30, 10])
FMS_BASE = [
    ("Pollen 외 (2021), Clin J Sport Med", "고교 선수", NA, "FMS 표준", "평균", 14.1, NA, "36개 연구 종합", "pollen"),
    ("Pollen 외 (2021)", "대학 선수", NA, "FMS 표준", "평균", 14.8, NA, "36개 연구 종합", "pollen"),
    ("Pollen 외 (2021)", "프로 선수", NA, "FMS 표준", "평균", 15.7, NA, "36개 연구 종합", "pollen"),
    ("Schneiders 외 (2011)", "젊은 활동 인구", NA, "FMS 표준", "평균", 15.7, NA, "95% CI 15.4–15.9", "schneiders"),
    ("게일릭 스포츠 남자 선수 (2013)", "남자 선수", NA, "FMS 표준", "평균", 15.56, 1.46, "", "gaelic"),
    ("청소년 학생 기준값", "청소년", NA, "FMS 표준", "평균", 14.59, NA, "95% CI 14.43–14.74", "adol"),
    ("대학생 100명", "대학생", 100, "FMS 표준", "평균", 14.40, NA, "", "college"),
]
plain_rows(ws, FMS_BASE)

# ------------------------------------------------------------------ 5. FMS 하위항목
ws = wb.create_sheet("FMS하위항목")
title(ws, "FMS 하위항목 사전/사후 (각 0–3점)", "선수 대상 항목별 사전·사후 자료는 드뭅니다. 아래는 확인된 값만 적었고, 대부분 비선수 자료입니다.")
pp_sheet(ws, [
    ("Frost 외 (2015) — Deep Squat", "소방관", NA, "채점 기준 안내 후 재검사", "FMS 표준", "단일군", 1.4, 0.7, 2.0, 0.6, "", "학습효과 연구", "frost"),
    ("Frost 외 (2015) — Hurdle Step", "소방관", NA, "채점 기준 안내 후 재검사", "FMS 표준", "단일군", 2.1, 0.4, 2.4, 0.5, "", "학습효과 연구", "frost"),
    ("Frost 외 (2015) — In-line Lunge", "소방관", NA, "채점 기준 안내 후 재검사", "FMS 표준", "단일군", 2.1, 0.4, 2.7, 0.5, "", "학습효과 연구", "frost"),
    ("Frost 외 (2015) — Rotary Stability", "소방관", NA, "채점 기준 안내 후 재검사", "FMS 표준", "단일군", 2.3, 0.5, 2.4, 0.5, "", "학습효과 연구", "frost"),
    ("고교 선수 기준값 — Trunk Stability Push-up", "고교 선수", NA, "횡단", "FMS 표준", "집단 A", 1.20, 0.49, None, None, "", "두 집단의 구분 원문 확인. 사전값 참고용", "hs"),
    ("고교 선수 기준값 — Trunk Stability Push-up", "고교 선수", NA, "횡단", "FMS 표준", "집단 B", 1.41, 0.72, None, None, "", "두 집단의 구분 원문 확인", "hs"),
    ("고교 선수 기준값 — Rotary Stability", "고교 선수", NA, "횡단", "FMS 표준", "집단 A", 1.51, 0.53, None, None, "", "두 집단의 구분 원문 확인", "hs"),
    ("고교 선수 기준값 — Rotary Stability", "고교 선수", NA, "횡단", "FMS 표준", "집단 B", 1.56, 0.55, None, None, "", "두 집단의 구분 원문 확인", "hs"),
    ("이진 외 (2017) — Deep Squat", "일반 성인", 18, "횡단", "FMS 표준", "전체", 2.61, NA, None, None, "", "비선수", "leejin"),
    ("이진 외 (2017) — Hurdle Step(우)", "일반 성인", 18, "횡단", "FMS 표준", "전체", 2.67, NA, None, None, "", "비선수", "leejin"),
])

# ------------------------------------------------------------------ 6. 메타분석
ws = wb.create_sheet("메타분석")
title(ws, "메타분석·체계적 고찰의 평균 효과", "평균 차이(MD)는 원 단위(mL/kg/min 또는 점), SMD는 표준화 효과크기입니다.")
header(ws, 4, ["연구", "변인", "비교", "효과 지표", "효과", "95% CI 하한", "95% CI 상한", "비고", "출처"],
       [30, 12, 26, 10, 9, 10, 10, 34, 10])
META = [
    ("Vasconcelos 외 (2019), J Strength Cond Res", "VO2max", "투기종목(타격계) HIIT", "MD", 2.83, 0.40, 5.25,
     "다른 투기종목 군 2.36. 개별 연구 향상폭 4.6–16.3%. 단위 표기 원문 확인", "vasc"),
    ("무술 선수 HIIT RCT 메타분석 (2026)", "VO2max", "HIIT vs 대조", "SMD", 1.04, 0.74, 1.33, "I² = 16.3%", "ma26"),
    ("Ma 외 (2023), Heliyon", "VO2max", "엘리트 선수 HIIT vs 기존 훈련", "SMD", 0.58, 0.30, 0.87, "9개 연구, 176명", "ma23"),
    ("Maleki 외 (2025), Sci Rep", "FMS 총점", "코어 안정성 훈련", "MD", 2.89, 2.46, 3.32, "40개 연구, 선수 1,604명", "maleki"),
    ("Maleki 외 (2025)", "FMS 총점", "신경근 훈련", "MD", 2.31, 1.64, 2.98, "", "maleki"),
    ("Maleki 외 (2025)", "FMS 총점", "저항 훈련", "MD", 2.21, 1.82, 2.61, "", "maleki"),
    ("Maleki 외 (2025)", "FMS 총점", "통합 훈련", "MD", 1.92, 1.57, 2.27, "", "maleki"),
    ("Maleki 외 (2025)", "FMS 총점", "기능성 훈련", "MD", 1.74, 1.34, 2.13, "이 연구의 H-IFIT와 가장 가까운 유형", "maleki"),
    ("Maleki 외 (2025)", "FMS 총점", "FIFA 11+ 준비운동", "MD", 1.69, 0.79, 2.60, "이질성 높음", "maleki"),
    ("고위험 선수 체계적 고찰 (2022)", "FMS 총점", "각종 중재(6개 연구)", "평균 증가", 3.28, None, None, "주 1–4회, 6–8주", "highrisk"),
    ("Teyhen 외 (2012), JOSPT", "FMS 총점", "측정오차(최소탐지변화량)", "MDC", 2.1, None, None, "평가자 간 2.1점, 평가자 내 2.5점", "teyhen"),
]
r = 5
for *vals, key in META:
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = F(size=9), BOX
        c.alignment = CEN if j in (4, 5, 6, 7) else WRAP
        if j in (5, 6, 7) and isinstance(v, (int, float)):
            c.number_format = "0.00"
    c = ws.cell(row=r, column=9, value="링크")
    c.hyperlink, c.font, c.border, c.alignment = URL[key], F(size=9, color="0563C1", underline="single"), BOX, CEN
    ws.row_dimensions[r].height = 30
    r += 1
ws.freeze_panes = "B5"

# ------------------------------------------------------------------ 7. 참고범위 요약
ws = wb.create_sheet("참고범위_요약")
title(ws, "참고범위 요약 (대학 겨루기 선수, 12주 중재 기준으로 선행연구를 종합한 범위)",
      "위 시트들의 값을 종합해 정리한 참고 범위입니다. 개별 연구 값이 아니며, 하위항목은 근거가 약합니다. 재분석 결과의 타당성 확인과 논의 비교에만 쓰세요.")
header(ws, 4, ["변인", "집단", "사전 하한", "사전 상한", "사후 하한", "사후 상한", "변화량 하한", "변화량 상한", "주요 근거", "근거 수준"],
       [22, 10, 9, 9, 9, 9, 10, 10, 46, 10])
SUM = [
    ("VO2max (mL/kg/min)", "실험군", 48, 55, 51, 58, 2, 4, "Batra & Zatoń 50.13→53.16; 투기종목 메타 MD 2.83; Xu +4.72% (+4–8%)", "중간"),
    ("VO2max (mL/kg/min)", "통제군", 48, 55, 48, 56, 0, 1, "Batra & Zatoń 통제군 변화 없음; Xu 통제군 +0.47%", "중간"),
    ("FMS 총점 (점)", "실험군", 13, 16, 16, 18, 1.7, 3.0, "Maleki MD 1.74–2.89; PeerJ 15.73→18.07; Kiesel 13.3→16.3", "중간"),
    ("FMS 총점 (점)", "통제군", 13, 16, 13, 17, 0, 1, "학습효과로 소폭 증가 가능(Kiesel 2011)", "약함"),
    ("Deep Squat", "실험군", 1.4, 2.0, None, None, 0.3, 0.6, "Frost 1.4→2.0", "약함"),
    ("Hurdle Step", "실험군", 2.0, 2.1, None, None, 0.2, 0.3, "Frost 2.1→2.4", "약함"),
    ("In-line Lunge", "실험군", 2.1, 2.4, None, None, 0.3, 0.6, "Frost 2.1→2.7", "약함"),
    ("Shoulder Mobility", "실험군", 2.0, 2.6, None, None, 0, 0, "프로그램에 어깨 가동성 운동 없음. 선수 수치 미확인", "매우 약함"),
    ("Active Straight Leg Raise", "실험군", 2.0, 2.5, None, None, 0, 0.2, "프로그램에 햄스트링 유연성 운동 없음. 선수 수치 미확인", "매우 약함"),
    ("Trunk Stability Push-up", "실험군", 1.2, 2.2, None, None, 0, 0.3, "고교 선수 1.20–1.41", "약함"),
    ("Rotary Stability", "실험군", 1.5, 2.3, None, None, 0.1, 0.3, "고교 선수 1.51–1.56; Frost 2.3→2.4", "약함"),
]
r = 5
for vals in SUM:
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = F(size=9), BOX
        c.alignment = WRAP if j in (1, 9) else CEN
        if 3 <= j <= 8 and isinstance(v, (int, float)):
            c.number_format = "0.0" if (isinstance(v, float) and v != int(v)) or vals[0] not in ("VO2max (mL/kg/min)",) else "0"
    # 사후 범위가 비어 있으면 사전+변화량으로 계산
    if vals[4] is None:
        ws.cell(row=r, column=5, value=f"=C{r}+G{r}").fill = CALC
        ws.cell(row=r, column=6, value=f"=MIN(3,D{r}+H{r})").fill = CALC
        for j in (5, 6):
            ws.cell(row=r, column=j).number_format = "0.0"
            ws.cell(row=r, column=j).font = F(size=9)
            ws.cell(row=r, column=j).alignment = CEN
    ws.row_dimensions[r].height = 30
    r += 1
ws.cell(row=r + 1, column=1, value=(
    "판단 기준: 실험군 VO2max 변화가 +10 이상, FMS 총점 변화가 +4 이상이면 선행연구 범위를 크게 벗어나므로 측정 절차와 입력값을 먼저 확인하세요. "
    "FMS 평균 변화가 2점 안팎이면 최소탐지변화량(2.1–2.5점)과 함께 해석하세요.")).font = F(size=9, color="595959")
ws.freeze_panes = "C5"

for w in wb.worksheets:
    w.sheet_view.zoomScale = 90
wb.save(OUT)
print("saved", OUT)
