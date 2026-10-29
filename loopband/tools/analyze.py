"""루프밴드 연구 원자료 → 원고 표와 자료 점검 결과.

사용: python3 analyze.py <원자료_입력양식_SPSS.xlsx 또는 .sav> <결과.xlsx>

계산 방식(SPSS 25 기본값과 같게 맞춤)
- 사전 동질성: 독립표본 t(Levene p<.05이면 Welch)
- 집단 내 비교: 사전-사후 차이값의 Shapiro-Wilk p≥.05이면 대응 t(사전-사후), 아니면 Wilcoxon 부호순위
  (동순위 보정, 연속성 보정 없음, 양측 점근 p, SPSS처럼 Z는 항상 음수로 표기)
- 반복측정 이원분산분석(2집단×2시기): 제Ⅲ유형. 시기 효과는 집단별 평균의 비가중 평균으로 검정
- 공분산분석(민감도 분석): 사후값 = 사전값 + 집단. 기울기 동질성(사전값×집단) 함께 보고
"""
import sys
import math
import numpy as np
from scipy import stats
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC, OUT = sys.argv[1], sys.argv[2]
VARS = [  # (열 접두어, 원고 표기, 단위, 값이 클수록 좋은가, 정수 기록 항목인가, 허용 범위)
    ("KFES", "낙상효능감(KFES-I)", "점", False, True, (16, 64)),
    ("의자일어서기", "의자 일어서기", "회", True, True, (0, 40)),
    ("팔굽히기", "팔 굽히기", "회", True, True, (0, 50)),
    ("등뒤로손잡기", "등 뒤로 손잡기", "cm", True, False, (-60, 30)),
    ("앉아윗몸앞으로굽히기", "의자에 앉아 윗몸 앞으로 굽히기", "cm", True, False, (-40, 30)),
    ("2분제자리걷기", "2분 제자리 걷기", "회", True, True, (0, 200)),
    ("2.44m갔다오기", "2.44 m 갔다 오기", "초", False, False, (2, 60)),
]
ALPHA_WITHIN = .025


def fp(p):
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return "-"
    return "<.001" if p < .001 else f"{p:.3f}".lstrip("0") if p < 1 else "1.000"


def msd(x):
    return f"{np.mean(x):.2f}±{np.std(x, ddof=1):.2f}"


# ------------------------------------------------------------------ 읽기
# 입력 양식(원자료_입력양식_SPSS.xlsx의 data 시트) 또는 SPSS .sav. 변수명은 build_spss_template.py 기준
def num(v):
    if v is None or v == "" or (isinstance(v, float) and math.isnan(v)):
        return None
    return float(v)


if SRC.lower().endswith(".sav"):
    import pyreadstat
    df, _ = pyreadstat.read_sav(SRC)
    raw = df.to_dict("records")
else:
    ws = load_workbook(SRC, data_only=True)["data"]
    heads = [c.value for c in ws[1]]
    raw = [dict(zip(heads, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
raw = [r for r in raw if num(r.get("id")) is not None]
if not raw:
    sys.exit("data 시트에 입력된 행이 없습니다.")

GROUP = {1: "실험", 2: "통제"}
RSN = {1: "출석 미달", 2: "건강상 이유", 3: "사후검사 불참", 4: "참여 철회", 5: "기타"}
WHEN = {1: "사전검사 후 중재 시작 전", 2: "중재 중", 3: "사후검사 전"}
kfes_partial = []


def best(r, stem, t, fn):
    v = [num(r.get(f"{stem}{k}_{t}")) for k in (1, 2)]
    v = [x for x in v if x is not None]
    return fn(v) if v else None


rows = []
for r in raw:
    code = lambda k: None if num(r.get(k)) is None else int(num(r.get(k)))
    d = {"ID": code("id"), "집단": GROUP.get(code("group")), "최종분석포함": "예" if code("incl") == 1 else "아니오",
         "제외사유": RSN.get(code("excl_rsn")) if code("excl_rsn") else None, "제외시점": WHEN.get(code("excl_when")),
         "출석횟수": r.get("attend"), "나이": r.get("age"), "신장": r.get("height"), "체중": r.get("weight")}
    if r.get("excl_txt"):
        d["제외사유"] = f"{d['제외사유'] or '사유 코드 없음'}({r['excl_txt']})"
    for t, tk in (("pre", "사전"), ("post", "사후")):
        items = [num(r.get(f"kfes{i:02d}_{t}")) for i in range(1, 17)]
        k = sum(x is not None for x in items)
        d[f"KFES_{tk}"] = sum(items) if k == 16 else None
        if 0 < k < 16:
            kfes_partial.append((d["ID"], tk, k))
        d[f"의자일어서기_{tk}"] = r.get(f"chair_{t}")
        d[f"팔굽히기_{tk}"] = r.get(f"curl_{t}")
        d[f"2분제자리걷기_{tk}"] = r.get(f"step_{t}")
        d[f"등뒤로손잡기_{tk}"] = best(r, "bs", t, max)
        d[f"앉아윗몸앞으로굽히기_{tk}"] = best(r, "sr", t, max)
        d[f"2.44m갔다오기_{tk}"] = best(r, "tug", t, min)
    rows.append(d)


def cronbach(t):
    m = np.array([[num(r.get(f"kfes{i:02d}_{t}")) for i in range(1, 17)] for r in raw
                  if num(r.get("incl")) == 1], dtype=float)
    m = m[~np.isnan(m).any(axis=1)]
    if len(m) < 3:
        return float("nan"), len(m)
    k = m.shape[1]
    return k / (k - 1) * (1 - m.var(axis=0, ddof=1).sum() / m.sum(axis=1).var(ddof=1)), len(m)


all_rows = rows
final = [r for r in rows if str(r.get("최종분석포함", "")).strip() == "예"]
grp = {g: [r for r in final if r["집단"] == g] for g in ("실험", "통제")}
nE, nC = len(grp["실험"]), len(grp["통제"])
N = nE + nC
if nE < 3 or nC < 3:
    sys.exit(f"최종 분석 대상이 너무 적습니다(실험 {nE}, 통제 {nC}).")


def col(g, name):
    return np.array([num(r[name]) for r in grp[g]], dtype=float)


def paired(g, v):
    pre, post = col(g, v + "_사전"), col(g, v + "_사후")
    ok = ~(np.isnan(pre) | np.isnan(post))
    return pre[ok], post[ok], int((~ok).sum())


# ------------------------------------------------------------------ 검정 함수
def indep(a, b):
    lev = stats.levene(a, b, center="mean").pvalue
    welch = lev < .05
    t = stats.ttest_ind(a, b, equal_var=not welch)
    return t.statistic, t.pvalue, lev, welch


def wilcoxon_spss(pre, post):
    d = post - pre
    nz = d[d != 0]
    k = len(nz)
    if k == 0:
        return 0.0, 1.0, k
    ranks = stats.rankdata(np.abs(nz))
    tp, tm = ranks[nz > 0].sum(), ranks[nz < 0].sum()
    mu = k * (k + 1) / 4
    _, cnt = np.unique(np.abs(nz), return_counts=True)
    var = k * (k + 1) * (2 * k + 1) / 24 - (cnt ** 3 - cnt).sum() / 48
    z = (min(tp, tm) - mu) / math.sqrt(var)
    return z, 2 * stats.norm.sf(abs(z)), k


def within(pre, post):
    d = pre - post
    n = len(d)
    if np.std(d, ddof=1) == 0:
        sw_p = float("nan")
        use_t = False
    else:
        sw_p = stats.shapiro(d).pvalue
        use_t = sw_p >= .05
    if use_t:
        t = stats.ttest_rel(pre, post)
        return dict(kind="t", stat=t.statistic, p=t.pvalue, df=n - 1, sw_p=sw_p,
                    dz=np.mean(d) / np.std(d, ddof=1))
    z, p, k = wilcoxon_spss(pre, post)
    sd = np.std(d, ddof=1)
    return dict(kind="Z", stat=z, p=p, df=None, sw_p=sw_p, nonzero=k,
                dz=(np.mean(d) / sd) if sd > 0 else float("nan"))


def mixed_anova(preE, postE, preC, postC):
    dE, dC = postE - preE, postC - preC
    nE_, nC_ = len(dE), len(dC)
    dfe = nE_ + nC_ - 2
    h = 1 / nE_ + 1 / nC_
    sp_d = ((nE_ - 1) * dE.var(ddof=1) + (nC_ - 1) * dC.var(ddof=1)) / dfe
    sE, sC = (preE + postE) / 2, (preC + postC) / 2
    sp_s = ((nE_ - 1) * sE.var(ddof=1) + (nC_ - 1) * sC.var(ddof=1)) / dfe
    out = {}
    out["Group"] = (sE.mean() - sC.mean()) ** 2 / (sp_s * h)
    out["Time"] = ((dE.mean() + dC.mean()) / 2) ** 2 / (sp_d * h / 4)
    out["GxT"] = (dE.mean() - dC.mean()) ** 2 / (sp_d * h)
    return {k: (F, stats.f.sf(F, 1, dfe), F / (F + dfe)) for k, F in out.items()}, dfe


def ols(X, y):
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)
    res = y - X @ beta
    dfe = len(y) - X.shape[1]
    mse = res @ res / dfe
    cov = mse * np.linalg.inv(X.T @ X)
    return beta, np.sqrt(np.diag(cov)), dfe, res @ res


def ancova(preE, postE, preC, postC):
    pre = np.r_[preE, preC]
    post = np.r_[postE, postC]
    g = np.r_[np.ones(len(preE)), np.zeros(len(preC))]
    X = np.c_[np.ones_like(pre), pre, g]
    b, se, dfe, sse = ols(X, post)
    F = (b[2] / se[2]) ** 2
    Xi = np.c_[X, pre * g]
    bi, sei, dfi, _ = ols(Xi, post)
    p_slope = 2 * stats.t.sf(abs(bi[3] / sei[3]), dfi)
    m = pre.mean()
    adjE, adjC = b[0] + b[1] * m + b[2], b[0] + b[1] * m
    overlap = (max(preE.min(), preC.min()), min(preE.max(), preC.max()))
    return dict(F=F, p=stats.f.sf(F, 1, dfe), df=dfe, eta=F / (F + dfe), adjE=adjE, adjC=adjC,
                se_diff=se[2], p_slope=p_slope, overlap=overlap)


# ------------------------------------------------------------------ 계산
out_rows = {}

# 대상자 흐름
flow = [["구분", "실험집단", "통제집단", "합계"]]
cnt = lambda cond, g: sum(1 for r in all_rows if r["집단"] == g and cond(r))
flow.append(["최초 참여", cnt(lambda r: True, "실험"), cnt(lambda r: True, "통제"), len(all_rows)])
reasons = sorted({(r.get("제외시점") or "시점 미기재", r.get("제외사유") or "사유 미기재")
                  for r in all_rows if str(r.get("최종분석포함", "")).strip() != "예"})
for when, why in reasons:
    f = lambda r: str(r.get("최종분석포함", "")).strip() != "예" and (r.get("제외시점") or "시점 미기재") == when \
        and (r.get("제외사유") or "사유 미기재") == why
    e, c = cnt(f, "실험"), cnt(f, "통제")
    flow.append([f"제외: {when} / {why}", e, c, e + c])
flow.append(["최종 분석", nE, nC, N])
att = [num(r.get("출석횟수")) for r in grp["실험"] if num(r.get("출석횟수")) is not None]
if att:
    flow.append(["실험집단 출석(최종 분석 대상)", f"{np.mean(att):.1f}±{np.std(att, ddof=1):.1f}회, 최소 {min(att):.0f}회", "", ""])
out_rows["대상자흐름"] = flow

# 표 1 + 사전 동질성
t1 = [["변인", f"실험집단(n={nE})", f"통제집단(n={nC})", "t", "p", "검정"]]
for v, lab in (("나이", "나이(yr)"), ("신장", "키(cm)"), ("체중", "체중(kg)")):
    a, b = col("실험", v), col("통제", v)
    a, b = a[~np.isnan(a)], b[~np.isnan(b)]
    t, p, lev, w = indep(a, b)
    t1.append([lab, msd(a), msd(b), f"{t:.3f}", fp(p), "Welch" if w else "Student"])
for v, lab, unit, *_ in VARS:
    a, b = col("실험", v + "_사전"), col("통제", v + "_사전")
    a, b = a[~np.isnan(a)], b[~np.isnan(b)]
    t, p, lev, w = indep(a, b)
    sp = math.sqrt(((len(a) - 1) * a.var(ddof=1) + (len(b) - 1) * b.var(ddof=1)) / (len(a) + len(b) - 2))
    t1.append([f"{lab} 사전({unit})", msd(a), msd(b), f"{t:.3f}", fp(p),
               ("Welch" if w else "Student") + f", d={(a.mean() - b.mean()) / sp:.2f}"])
out_rows["표1_동질성"] = t1

# 결과 표
res = [["변인", "집단", "PRE M±SD", "POST M±SD", "집단 내 검정", "p", "Bonferroni(α=.025)", "효과크기 dz",
        "효과", "F(1,%d)" % (N - 2), "p", "ηp²"]]
anc = [["변인", "조정 평균 실험", "조정 평균 통제", "F", "df", "p", "ηp²", "기울기 동질성 p", "사전값 겹침 구간", "판단"]]
norm = [["변인", "집단", "차이값 Shapiro-W p", "적용 검정", "변화 0인 인원"]]
for v, lab, unit, higher_better, *_ in VARS:
    preE, postE, mE = paired("실험", v)
    preC, postC, mC = paired("통제", v)
    aov, dfe = mixed_anova(preE, postE, preC, postC)
    eff = ["Group", "Time", "GxT"]
    for gi, (g, pre, post) in enumerate((("실험", preE, postE), ("통제", preC, postC))):
        w = within(pre, post)
        stat = f"t({w['df']})={w['stat']:.3f}" if w["kind"] == "t" else f"Z={w['stat']:.3f}"
        sig = "유의" if w["p"] < ALPHA_WITHIN else "유의하지 않음"
        res.append([f"{lab}({unit})" if gi == 0 else "", g, msd(pre), msd(post), stat, fp(w["p"]), sig,
                    f"{w['dz']:.2f}" if not math.isnan(w["dz"]) else "-",
                    "", "", "", ""])
        norm.append([lab, g, fp(w["sw_p"]), "대응 t" if w["kind"] == "t" else "Wilcoxon",
                     int((post - pre == 0).sum())])
    base = len(res) - 2
    for k, e in enumerate(eff):
        F, p, eta = aov[e]
        if k < 2:
            res[base + k][8:12] = [e, f"{F:.3f}", fp(p), f"{eta:.2f}"]
        else:
            res.append(["", "", "", "", "", "", "", "", e, f"{F:.3f}", fp(p), f"{eta:.2f}"])
    a = ancova(preE, postE, preC, postC)
    lo, hi = a["overlap"]
    judge = []
    if a["p_slope"] < .05:
        judge.append("기울기 동질성 위배: ANCOVA 해석 제한")
    if lo > hi:
        judge.append("두 집단 사전값이 겹치지 않음: 조정 평균은 외삽값")
    anc.append([lab, f"{a['adjE']:.2f}", f"{a['adjC']:.2f}", f"{a['F']:.3f}", f"1, {a['df']}", fp(a["p"]),
                f"{a['eta']:.2f}", fp(a["p_slope"]), f"{lo:g}~{hi:g}" if lo <= hi else "겹침 없음",
                "; ".join(judge) or "-"])
out_rows["결과표"] = res
out_rows["공분산분석"] = anc
out_rows["정규성"] = norm

# 자료 점검
chk = [["점검", "변인", "집단", "결과", "의미"]]
for v, lab, unit, _, integer, (lo, hi) in VARS:
    for g in ("실험", "통제"):
        pre, post, miss = paired(g, v)
        same = int((pre == post).sum())
        chk.append(["사전=사후 동일 인원", lab, g, f"{same}/{len(pre)}명",
                    "절반 이상이면 기록지에서 사후값이 실제로 측정된 값인지 확인" if same * 2 >= len(pre) else ""])
        if miss:
            chk.append(["결측", lab, g, f"{miss}명", "최종 분석 대상인데 값이 비어 있음"])
        vals = np.r_[pre, post]
        frac = vals[np.abs(vals - np.round(vals)) > 1e-9]
        if integer and len(frac):
            chk.append(["정수가 아닌 값", lab, g, f"{len(frac)}개", "횟수·총점 항목은 정수여야 함"])
        if v == "2.44m갔다오기" and len(vals) and not len(frac):
            chk.append(["기록 단위", lab, g, "모두 정수 초", "SFT 기준은 0.1초 단위. 기록지 단위를 원고에 명시"])
        out_of = vals[(vals < lo) | (vals > hi)]
        if len(out_of):
            chk.append(["범위 밖 값", lab, g, ", ".join(f"{x:g}" for x in out_of), f"허용 범위 {lo}~{hi}"])
dup = {}
for r in final:
    key = tuple(r.get(v + s) for v, *_ in VARS for s in ("_사전", "_사후"))
    dup.setdefault(key, []).append(r["ID"])
for ids in dup.values():
    if len(ids) > 1:
        chk.append(["모든 측정값이 같은 대상자", "-", "-", ", ".join(map(str, ids)), "중복 입력 여부 확인"])
for i, tk, k in kfes_partial:
    chk.append(["KFES 문항 일부 결측", "낙상효능감(KFES-I)", tk, f"ID {i}: 16문항 중 {k}문항", "총점을 계산하지 않음(결측). 설문지 확인"])
out_rows["자료점검"] = chk
out_rows["신뢰도"] = [["척도", "Cronbach α", "사용 인원"]] + [
    [f"KFES-I {tk}", f"{a:.3f}" if not math.isnan(a) else "-", n] for t, tk in (("pre", "사전"), ("post", "사후"))
    for a, n in [cronbach(t)]]

# ------------------------------------------------------------------ 쓰기
FONT = "맑은 고딕"
hdr_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor="1F4E78")
base_font = Font(name=FONT, size=10)
side = Side(style="thin", color="BFBFBF")
box = Border(left=side, right=side, top=side, bottom=side)
wb = Workbook()
wb.remove(wb.active)
notes = {
    "결과표": "집단 내 검정: 차이값 정규성에 따라 대응 t(사전-사후) 또는 Wilcoxon(Z, SPSS 표기처럼 음수). "
             "반복측정 이원분산분석은 제Ⅲ유형. 원고 표에는 t와 Z를 구분해 적을 것.",
    "공분산분석": "사전값이 겹치지 않는 변인은 조정 평균이 관측 범위 밖의 외삽값이므로 공분산분석으로도 사전 차이를 보정했다고 쓸 수 없음.",
}
for name, table in out_rows.items():
    sh = wb.create_sheet(name)
    for i, row in enumerate(table, 1):
        for j, v in enumerate(row, 1):
            c = sh.cell(i, j, v)
            c.border = box
            c.font = hdr_font if i == 1 else base_font
            if i == 1:
                c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for j in range(1, len(table[0]) + 1):
        sh.column_dimensions[chr(64 + j)].width = 22 if j == 1 else 14
    if name in notes:
        c = sh.cell(len(table) + 2, 1, notes[name])
        c.font = Font(name=FONT, size=9, color="595959")
wb.save(OUT)
print(f"saved {OUT}  (최종 분석 실험 {nE}명, 통제 {nC}명)")
