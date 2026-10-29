"""루프밴드 연구 SPSS 원자료 입력 양식 생성.

사용: python3 build_spss_template.py <출력 폴더>

만드는 파일(변수 정의는 아래 VARS 하나에서 나오므로 세 파일이 항상 일치한다)
- 원자료_입력양식_SPSS.xlsx : 'data' 시트 1행이 SPSS 변수명. 코드북·입력안내·입력점검 시트 포함
- 원자료_SPSS_빈파일.sav    : 변수명·라벨·값 라벨·측정수준이 정의된 빈 SPSS 파일(SPSS에서 바로 입력할 때)
- 루프밴드_분석.sps          : 엑셀 불러오기, 라벨, 파생변수, 입력 점검, 원고 표에 필요한 분석 구문
"""
import os
import sys
import pandas as pd
import pyreadstat
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L

OUTDIR = sys.argv[1]
XLSX = "원자료_입력양식_SPSS.xlsx"
SAV = "원자료_SPSS_빈파일.sav"
SPS = "루프밴드_분석.sps"
NROWS = 40          # 입력 가능 행(30명 + 여유)

# ------------------------------------------------------------------ 변수 정의
# (변수명, 라벨, 종류, 소수자리, 범위 또는 코드, 기록 방법, 구역)
# 종류: code(값 라벨), int(정수), dec(소수), str(문자), date(날짜)
CODES = {
    "group": {1: "실험집단", 2: "통제집단"},
    "floor": {2: "2층", 4: "4층"},
    "incl": {1: "최종 분석 포함", 0: "제외"},
    "excl_rsn": {1: "출석 미달", 2: "건강상 이유", 3: "사후검사 불참", 4: "참여 철회", 5: "기타"},
    "excl_when": {1: "사전검사 후 중재 시작 전", 2: "중재 중", 3: "사후검사 전"},
}
VARS = [
    ("id", "대상자 번호(기록지 번호)", "int", 0, (1, 99), "기록지에 적힌 번호", "기본"),
    ("group", "집단", "code", 0, CODES["group"], "1=실험, 2=통제", "기본"),
    ("floor", "이용 층", "code", 0, CODES["floor"], "2 또는 4", "기본"),
    ("incl", "최종 분석 포함", "code", 0, CODES["incl"], "1=포함, 0=제외", "기본"),
    ("excl_rsn", "제외 사유", "code", 0, CODES["excl_rsn"], "제외자만. 1 출석 미달, 2 건강상 이유, 3 사후검사 불참, 4 참여 철회, 5 기타", "기본"),
    ("excl_txt", "제외 사유 상세", "str", 0, None, "제외자만. 예: 무릎 통증으로 4주차부터 불참", "기본"),
    ("excl_when", "제외 시점", "code", 0, CODES["excl_when"], "제외자만. 1 사전검사 후 중재 시작 전, 2 중재 중, 3 사후검사 전", "기본"),
    ("attend", "출석 횟수(24회 중)", "int", 0, (0, 24), "실험집단만. 출석부 기준", "기본"),
    ("band_fin", "8주차 최종 밴드 색상", "str", 0, None, "실험집단만. 예: 빨강", "기본"),
    ("age", "나이(만, 세)", "int", 0, (60, 105), "만 나이", "기본"),
    ("height", "신장(cm)", "dec", 1, (120, 190), "cm, 기록지 그대로", "기본"),
    ("weight", "체중(kg)", "dec", 1, (25, 120), "kg, 기록지 그대로", "기본"),
]
for t, tk in (("pre", "사전"), ("post", "사후")):
    VARS += [
        (f"date_{t}", f"{tk}검사 날짜", "date", 0, None, "YYYY-MM-DD", tk),
        (f"rater_{t}", f"{tk} 측정자", "str", 0, None, "이름 또는 기호(예: A)", tk),
    ]
    VARS += [(f"kfes{i:02d}_{t}", f"KFES-I {i}번 문항 {tk}", "int", 0, (1, 4),
              "1 전혀 걱정 안 함 ~ 4 매우 걱정", tk) for i in range(1, 17)]
    VARS += [
        (f"chair_{t}", f"의자 일어서기 {tk}(회/30초)", "int", 0, (0, 40), "30초 동안 완전히 일어선 횟수", tk),
        (f"curl_{t}", f"팔 굽히기 {tk}(회/30초)", "int", 0, (0, 50), "30초 동안 덤벨 들기 횟수", tk),
        (f"step_{t}", f"2분 제자리 걷기 {tk}(회)", "int", 0, (0, 200), "2분 동안 오른쪽 무릎이 기준 높이에 닿은 횟수", tk),
        (f"bs1_{t}", f"등 뒤로 손잡기 {tk} 1회(cm)", "dec", 1, (-60, 30), "닿지 않으면 -, 겹치면 +", tk),
        (f"bs2_{t}", f"등 뒤로 손잡기 {tk} 2회(cm)", "dec", 1, (-60, 30), "2회 기록이 없으면 비움", tk),
        (f"sr1_{t}", f"의자에 앉아 윗몸 앞으로 굽히기 {tk} 1회(cm)", "dec", 1, (-40, 30), "발끝 못 미치면 -, 지나면 +", tk),
        (f"sr2_{t}", f"의자에 앉아 윗몸 앞으로 굽히기 {tk} 2회(cm)", "dec", 1, (-40, 30), "2회 기록이 없으면 비움", tk),
        (f"tug1_{t}", f"2.44 m 갔다 오기 {tk} 1회(초)", "dec", 2, (2, 60), "초, 기록지에 적힌 자리수 그대로", tk),
        (f"tug2_{t}", f"2.44 m 갔다 오기 {tk} 2회(초)", "dec", 2, (2, 60), "2회 기록이 없으면 비움", tk),
    ]
VARS.append(("note", "비고", "str", 0, None, "기록지 특이사항(수정 흔적, 판독 불가 등)", "기본"))
NAMES = [v[0] for v in VARS]
COL = {v[0]: L(i + 1) for i, v in enumerate(VARS)}

# 분석 변수: (파생변수명, 원고 표기, 사전 계산식, 사후 계산식)
def best(stem, fn):
    return f"{fn}({stem}1_pre, {stem}2_pre)", f"{fn}({stem}1_post, {stem}2_post)"
DERIVED = [
    ("kfes", "낙상효능감(KFES-I)", "SUM.16(kfes01_pre TO kfes16_pre)", "SUM.16(kfes01_post TO kfes16_post)"),
    ("chair", "의자 일어서기", "chair_pre", "chair_post"),
    ("curl", "팔 굽히기", "curl_pre", "curl_post"),
    ("bs", "등 뒤로 손잡기", *best("bs", "MAX")),
    ("sr", "의자에 앉아 윗몸 앞으로 굽히기", *best("sr", "MAX")),
    ("step", "2분 제자리 걷기", "step_pre", "step_post"),
    ("tug", "2.44 m 갔다 오기", *best("tug", "MIN")),
]

# ------------------------------------------------------------------ 엑셀
FONT = "맑은 고딕"
F = lambda **k: Font(name=FONT, size=k.pop("size", 10), **k)
FILL = {"기본": "FFF2CC", "사전": "DDEBF7", "사후": "E2EFDA"}
HDR = {"기본": "7F6000", "사전": "1F4E78", "사후": "375623"}
side = Side(style="thin", color="BFBFBF")
BOX = Border(left=side, right=side, top=side, bottom=side)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws = wb.active
ws.title = "data"
for c, (name, label, kind, dec, rng, rule, part) in enumerate(VARS, 1):
    h = ws.cell(1, c, name)
    h.font = F(bold=True, color="FFFFFF", size=9)
    h.fill = PatternFill("solid", fgColor=HDR[part])
    h.alignment, h.border = CENTER, BOX
    h.comment = Comment(f"{label}\n{rule}", "양식")
    ws.column_dimensions[L(c)].width = {"str": 16, "date": 11}.get(kind, 7 if name.startswith("kfes") else 9)
    for r in range(2, NROWS + 2):
        cell = ws.cell(r, c)
        cell.fill = PatternFill("solid", fgColor=FILL[part])
        cell.border, cell.font = BOX, F()
        cell.alignment = Alignment(horizontal="center")
        if kind == "date":
            cell.number_format = "yyyy-mm-dd"
        elif kind == "dec":
            cell.number_format = "0." + "0" * dec
    rng_cells = f"{L(c)}2:{L(c)}{NROWS + 1}"
    if kind == "code":
        codes = ",".join(str(k) for k in rng)
        dv = DataValidation(type="list", formula1=f'"{codes}"', allow_blank=True,
                            error=f"{label}: {rule}", errorTitle="입력값 확인")
    elif kind in ("int", "dec"):
        dv = DataValidation(type="whole" if kind == "int" else "decimal", operator="between",
                            formula1=str(rng[0]), formula2=str(rng[1]), allow_blank=True,
                            error=f"{label}: {rng[0]}~{rng[1]} 범위의 {'정수' if kind == 'int' else '숫자'}", errorTitle="입력값 확인")
    elif kind == "date":
        dv = DataValidation(type="date", operator="greaterThan", formula1="45000", allow_blank=True,
                            error="날짜 형식(예: 2026-03-02)으로 입력", errorTitle="입력값 확인")
    else:
        dv = None
    if dv:
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(rng_cells)
ws.freeze_panes = "B2"
ws.row_dimensions[1].height = 30

# 코드북
cb = wb.create_sheet("코드북")
cb.append(["변수명", "설명", "구역", "종류", "입력값·코드", "허용 범위", "기록 방법"])
for name, label, kind, dec, rng, rule, part in VARS:
    codes = "; ".join(f"{k}={v}" for k, v in rng.items()) if kind == "code" else ""
    rr = f"{rng[0]}~{rng[1]}" if kind in ("int", "dec") else ""
    kk = {"code": "코드", "int": "정수", "dec": f"숫자(소수 {dec}자리까지)", "str": "문자", "date": "날짜"}[kind]
    cb.append([name, label, part, kk, codes, rr, rule])
cb.append([])
cb.append(["분석 때 계산되는 변수(입력하지 않음)"])
for stem, lab, pre, post in DERIVED:
    cb.append([f"{stem}_pre / {stem}_post", lab, "", "계산", f"사전 = {pre}", f"사후 = {post}", ""])
for i, w in enumerate((18, 34, 6, 14, 40, 12, 40), 1):
    cb.column_dimensions[L(i)].width = w
for c in cb[1]:
    c.font, c.fill = F(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
for row in cb.iter_rows(min_row=2):
    for c in row:
        c.font, c.alignment = F(size=9), WRAP

# 입력 점검(엑셀이 열 때 계산). data 시트는 SPSS가 그대로 읽으므로 수식은 이 시트에만 둔다
ck = wb.create_sheet("입력점검")
heads = ["id", "group", "incl", "KFES 사전 문항 수", "KFES 사후 문항 수", "KFES 사전 총점", "KFES 사후 총점",
         "등뒤 사전 최고", "등뒤 사후 최고", "앉아 사전 최고", "앉아 사후 최고", "2.44 사전 최고", "2.44 사후 최고",
         "사전=사후 같은 항목 수(7개 중)", "포함 대상 빈칸 수"]
ck.append(heads)
need = [n for n in NAMES if n.split("_")[-1] in ("pre", "post") and not n.startswith(("date", "rater"))
        and not n[:3] in ("bs2", "sr2") and not n.startswith("tug2")] + ["age", "height", "weight"]
for r in range(2, NROWS + 2):
    d = lambda n: f"data!{COL[n]}{r}"
    rngk = lambda t: f"data!{COL['kfes01_' + t]}{r}:{COL['kfes16_' + t]}{r}"
    kpre = f'IF(COUNT({rngk("pre")})=16,SUM({rngk("pre")}),"")'
    kpost = f'IF(COUNT({rngk("post")})=16,SUM({rngk("post")}),"")'
    mx = lambda s, t, fn: f'IF(COUNT({d(s + "1_" + t)},{d(s + "2_" + t)})=0,"",{fn}({d(s + "1_" + t)},{d(s + "2_" + t)}))'
    pairs = [("F", "G"), (d("chair_pre"), d("chair_post")), (d("curl_pre"), d("curl_post")),
             (d("step_pre"), d("step_post")), ("H", "I"), ("J", "K"), ("L", "M")]
    same = "+".join(
        (f'IF(OR({a}{r}="",{b}{r}=""),0,--({a}{r}={b}{r}))' if len(a) == 1 else
         f'IF(OR({a}="",{b}=""),0,--({a}={b}))') for a, b in pairs)
    blanks = "+".join(f'--({d(n)}="")' for n in need)
    row = [f'=IF({d("id")}="","",{d("id")})', f'=IF({d("group")}="","",{d("group")})', f'=IF({d("incl")}="","",{d("incl")})',
           f'=IF({d("id")}="","",COUNT({rngk("pre")}))', f'=IF({d("id")}="","",COUNT({rngk("post")}))',
           f"=IF({d('id')}=\"\",\"\",{kpre})", f"=IF({d('id')}=\"\",\"\",{kpost})",
           "=" + mx("bs", "pre", "MAX"), "=" + mx("bs", "post", "MAX"), "=" + mx("sr", "pre", "MAX"),
           "=" + mx("sr", "post", "MAX"), "=" + mx("tug", "pre", "MIN"), "=" + mx("tug", "post", "MIN"),
           f'=IF({d("id")}="","",{same})', f'=IF({d("incl")}<>1,"",{blanks})']
    ck.append(row)
for i in range(1, len(heads) + 1):
    ck.column_dimensions[L(i)].width = 11
for c in ck[1]:
    c.font, c.alignment = F(bold=True, size=9, color="FFFFFF"), CENTER
    c.fill = PatternFill("solid", fgColor="1F4E78")
ck.row_dimensions[1].height = 45
ck.freeze_panes = "B2"
red = PatternFill("solid", fgColor="F8CBAD")
ck.conditional_formatting.add(f"D2:E{NROWS + 1}", CellIsRule(operator="between", formula=["1", "15"], fill=red))
ck.conditional_formatting.add(f"N2:N{NROWS + 1}", CellIsRule(operator="greaterThanOrEqual", formula=["4"], fill=red))
ck.conditional_formatting.add(f"O2:O{NROWS + 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=red))

gd = wb.create_sheet("입력안내")
guide = [
    ("입력 순서", True),
    ("1. 'data' 시트에만 입력합니다. 1행의 변수명은 SPSS가 그대로 읽으므로 바꾸거나 행을 끼워 넣지 마세요.", False),
    ("2. 한 사람이 한 행입니다. 최초 참여자 30명 전원(제외자 포함)을 기록지 번호 순서로 입력합니다.", False),
    ("3. 사전 기록지를 보고 노란·파란 칸(기본, 사전)을 먼저 채우고, 사후 기록지를 보고 초록 칸(사후)을 채웁니다.", False),
    ("4. 기록지에 없는 값은 비워 둡니다. 0이나 다른 시점 값을 넣지 않습니다(빈칸은 SPSS에서 결측).", False),
    ("5. 기록지에 적힌 숫자를 그대로 입력합니다. 반올림하거나 단위를 바꾸지 않습니다. 2.44 m가 정수 초로 적혀 있으면 정수로 입력합니다.", False),
    ("6. 유연성과 2.44 m는 기록지에 2회가 있으면 1회·2회 칸에 각각, 최종값 1개만 있으면 1회 칸에만 입력합니다. 최고값은 분석 때 계산됩니다.", False),
    ("7. KFES-I는 16문항을 각각 입력합니다. 총점은 분석 때 계산되며, 원래 기록지 총점과 다르면 비고에 적습니다.", False),
    ("8. 제외자는 incl=0, 제외 사유·상세·시점을 입력합니다. 통제집단 제외자도 실제 사유를 적습니다.", False),
    ("9. 각 칸에 마우스를 올리면 설명이 나옵니다. 범위 밖 값이나 정해진 코드가 아니면 입력이 막힙니다.", False),
    ("입력 확인", True),
    ("'입력점검' 시트에서 빨간 칸을 확인합니다: KFES 문항이 1~15개만 입력된 행, 사전=사후 같은 항목이 4개 이상인 행, 포함 대상인데 빈칸이 있는 행.", False),
    ("가능하면 다른 사람이 같은 양식에 한 번 더 입력하고 두 파일을 비교합니다(이중 입력). 심사자가 원자료 확인을 요구했으므로 입력 파일과 기록지는 함께 보관합니다.", False),
    ("SPSS로 불러오기", True),
    ("루프밴드_분석.sps를 SPSS 구문 창에서 열고, 맨 위 FILE= 경로를 이 파일 위치로 바꾼 뒤 전체 실행합니다. 라벨, 계산 변수, 입력 점검, 원고 표 분석이 차례로 실행됩니다.", False),
    ("SPSS에 직접 입력하고 싶다면 원자료_SPSS_빈파일.sav를 열어 데이터 보기에서 입력하고, 구문의 GET DATA 대신 GET FILE 줄을 사용합니다.", False),
]
gd.column_dimensions["A"].width = 120
for i, (text, bold) in enumerate(guide, 1):
    c = gd.cell(i, 1, text)
    c.font, c.alignment = (F(bold=True, size=11) if bold else F()), WRAP
wb.save(os.path.join(OUTDIR, XLSX))

# ------------------------------------------------------------------ 빈 .sav
cols = {}
for name, label, kind, dec, rng, rule, part in VARS:
    cols[name] = pd.Series([], dtype="object" if kind == "str" else ("datetime64[ns]" if kind == "date" else "float64"))
df = pd.DataFrame(cols)
fmt = {}
for name, label, kind, dec, rng, rule, part in VARS:
    fmt[name] = {"str": "A100", "date": "SDATE10", "code": "F1.0", "int": "F3.0", "dec": f"F6.{dec}"}[kind]
pyreadstat.write_sav(
    df, os.path.join(OUTDIR, SAV),
    column_labels={v[0]: v[1] for v in VARS},
    variable_value_labels={k: v for k, v in CODES.items()},
    variable_measure={v[0]: ("nominal" if v[2] in ("code", "str") else "scale") for v in VARS},
    variable_format=fmt,
    variable_display_width={v[0]: 10 for v in VARS},
)

# ------------------------------------------------------------------ .sps
def q(s):
    return "'" + s.replace("'", "''") + "'"


pre_list = " ".join(f"{s}_pre" for s, *_ in DERIVED)
post_list = " ".join(f"{s}_post" for s, *_ in DERIVED)
L_ = []
w = L_.append
w("* 루프밴드 연구: 원자료 불러오기, 변수 정의, 입력 점검, 원고 표 분석.")
w("* build_spss_template.py가 만든 구문. 변수를 바꾸려면 그 스크립트를 고쳐서 다시 만든다.")
w("* 실행 전: 아래 FILE= 경로를 실제 엑셀 파일 위치로 바꾼다.")
w("")
w("* ---- 1. 불러오기 (엑셀 양식에 입력한 경우).")
w("GET DATA /TYPE=XLSX")
w("  /FILE='C:\\루프밴드\\원자료_입력양식_SPSS.xlsx'")
w("  /SHEET=NAME 'data'")
w("  /CELLRANGE=FULL")
w("  /READNAMES=ON")
w("  /ASSUMEDSTRWIDTH=100.")
w("* 빈 .sav에 직접 입력한 경우에는 위 GET DATA 대신 다음 줄의 * 를 지우고 사용.")
w("* GET FILE='C:\\루프밴드\\원자료_SPSS_빈파일.sav'.")
w("SELECT IF NOT SYSMIS(id).")
w("EXECUTE.")
w("DATASET NAME raw WINDOW=FRONT.")
w("")
w("* ---- 2. 변수 정의.")
w("VARIABLE LABELS")
w("\n".join(f"  {n} {q(lab)}" for n, lab, *_ in VARS) + ".")
w("VALUE LABELS")
w("\n  /".join(f" {k} " + " ".join(f"{c} {q(v)}" for c, v in codes.items()) for k, codes in CODES.items()) + ".")
w("VARIABLE LEVEL " + " ".join(k for k in CODES) + " (NOMINAL).")
w("FORMATS " + " ".join(f"{v[0]} (F6.{v[3]})" for v in VARS if v[2] == "dec") + ".")
w("")
w("* ---- 3. 계산 변수: KFES-I 총점(16문항 모두 있을 때만), 유연성 최고값(MAX), 2.44 m 최고값(MIN).")
for stem, lab, pre, post in DERIVED:
    if pre != f"{stem}_pre":
        w(f"COMPUTE {stem}_pre = {pre}.")
        w(f"COMPUTE {stem}_post = {post}.")
for stem, lab, *_ in DERIVED:
    w(f"COMPUTE d_{stem} = {stem}_pre - {stem}_post.")
    w(f"COMPUTE same_{stem} = ({stem}_pre = {stem}_post).")
w("VARIABLE LABELS")
w("\n".join(f"  {s}_pre {q(lab + ' 사전')}\n  {s}_post {q(lab + ' 사후')}\n  d_{s} {q(lab + ' 차이(사전-사후)')}\n  same_{s} {q(lab + ' 사전=사후 동일(1)')}"
            for s, lab, *_ in DERIVED) + ".")
w("EXECUTE.")
w("")
w("* ---- 4. 입력 점검: 범위, 결측, 사전=사후 동일 인원. 원고 분석 전에 기록지와 대조.")
w("FREQUENCIES VARIABLES=group incl excl_rsn excl_when /ORDER=ANALYSIS.")
w("CROSSTABS /TABLES=excl_rsn excl_when BY group /CELLS=COUNT.")
w(f"DESCRIPTIVES VARIABLES=age height weight attend {pre_list} {post_list} /STATISTICS=MEAN STDDEV MIN MAX.")
w("MEANS TABLES=" + " ".join(f"same_{s}" for s, *_ in DERIVED) + " BY group /CELLS=SUM COUNT.")
w("LIST VARIABLES=id group incl " + " ".join(f"{s}_pre {s}_post" for s, *_ in DERIVED[:4]) + ".")
w("LIST VARIABLES=id group incl " + " ".join(f"{s}_pre {s}_post" for s, *_ in DERIVED[4:]) + ".")
w("")
w("* ---- 5. 최종 분석 대상만 사용(incl=1).")
w("USE ALL.")
w("COMPUTE filter_incl = (incl = 1).")
w("FILTER BY filter_incl.")
w("EXECUTE.")
w("")
w("* 5-1. 표 1과 사전 동질성: 독립표본 t(Levene 결과에 따라 등분산 가정 행 또는 Welch 행을 사용).")
w(f"T-TEST GROUPS=group(1 2) /VARIABLES=age height weight {pre_list} /MISSING=ANALYSIS /CRITERIA=CI(.95).")
w("")
w("* 5-2. KFES-I 신뢰도(Cronbach α).")
w("RELIABILITY /VARIABLES=kfes01_pre TO kfes16_pre /SCALE('KFES-I 사전') ALL /MODEL=ALPHA.")
w("RELIABILITY /VARIABLES=kfes01_post TO kfes16_post /SCALE('KFES-I 사후') ALL /MODEL=ALPHA.")
w("")
w("* 5-3. 집단별 차이값 정규성(Shapiro-Wilk). p≥.05면 대응 t, p<.05면 Wilcoxon 결과를 원고에 사용.")
w("EXAMINE VARIABLES=" + " ".join(f"d_{s}" for s, *_ in DERIVED) + " BY group")
w("  /PLOT NPPLOT /STATISTICS NONE /NOTOTAL /MISSING=PAIRWISE.")
w("")
w("* 5-4. 집단 내 사전·사후 비교(집단별). 대응 t는 사전-사후, Wilcoxon Z는 SPSS 출력 그대로 보고.")
w("SORT CASES BY group.")
w("SPLIT FILE LAYERED BY group.")
w(f"T-TEST PAIRS={pre_list} WITH {post_list} (PAIRED) /CRITERIA=CI(.95) /MISSING=ANALYSIS.")
w(f"NPAR TESTS /WILCOXON={pre_list} WITH {post_list} (PAIRED) /MISSING ANALYSIS.")
w("SPLIT FILE OFF.")
w("")
w("* 5-5. 반복측정 이원분산분석(집단×시기, 제Ⅲ유형, ηp²).")
for s, lab, *_ in DERIVED:
    w(f"* {lab}.")
    w(f"GLM {s}_pre {s}_post BY group")
    w("  /WSFACTOR=time 2 Polynomial")
    w("  /METHOD=SSTYPE(3)")
    w("  /PRINT=DESCRIPTIVE ETASQ")
    w("  /CRITERIA=ALPHA(.05)")
    w("  /WSDESIGN=time")
    w("  /DESIGN=group.")
w("")
w("* 5-6. 민감도 분석: 사전값을 공변량으로 한 공분산분석과 회귀기울기 동질성.")
w("* 기울기 동질성: 아래 첫 번째 분석에서 group*사전값 상호작용 p<.05면 공분산분석 해석 제한.")
for s, lab, *_ in DERIVED:
    w(f"* {lab}.")
    w(f"UNIANOVA {s}_post BY group WITH {s}_pre")
    w("  /METHOD=SSTYPE(3)")
    w(f"  /DESIGN={s}_pre group {s}_pre*group.")
    w(f"UNIANOVA {s}_post BY group WITH {s}_pre")
    w("  /METHOD=SSTYPE(3)")
    w(f"  /EMMEANS=TABLES(group) WITH({s}_pre=MEAN)")
    w("  /PRINT=DESCRIPTIVE ETASQ PARAMETER")
    w(f"  /DESIGN={s}_pre group.")
w("")
w("* 5-7. 사전값 분포가 겹치는지 확인(공분산분석 조정 평균 해석 전에 확인).")
w(f"MEANS TABLES={pre_list} BY group /CELLS=COUNT MEAN STDDEV MIN MAX.")
w("FILTER OFF.")
w("USE ALL.")
w("EXECUTE.")
with open(os.path.join(OUTDIR, SPS), "w", encoding="utf-8-sig", newline="\r\n") as f:
    f.write("\n".join(L_) + "\n")
print("saved", XLSX, SAV, SPS)
