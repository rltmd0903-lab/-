"""루프밴드 연구 원자료 입력 양식(.xlsx) 생성.

사용: python3 build_input_sheet.py <출력경로.xlsx>

최초 참여자 30명 전원을 측정 기록지 그대로 입력한다(제외자 포함).
분석은 analyze.py가 '최종분석포함=예'인 행만 사용한다.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
FONT = "맑은 고딕"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HDR = Font(name=FONT, size=9, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, size=9, color="595959")
FILL_HDR = PatternFill("solid", fgColor="1F4E78")
FILL_PRE = PatternFill("solid", fgColor="DDEBF7")
FILL_POST = PatternFill("solid", fgColor="E2EFDA")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(wrap_text=True, vertical="top")

# (열 이름, 너비, 단위·기록 방법)  열 이름은 analyze.py가 그대로 읽으므로 바꾸지 말 것
COLS = [
    ("ID", 7, "기록지 번호"),
    ("집단", 7, "실험 / 통제"),
    ("이용층", 7, "2층 / 4층"),
    ("최종분석포함", 8, "예 / 아니오"),
    ("제외사유", 16, "출석 미달, 건강상 이유(구체적으로), 사후검사 불참 등"),
    ("제외시점", 12, "사전검사 후 / 중재 중 / 사후검사 전"),
    ("출석횟수", 7, "24회 중 참여 횟수(통제집단은 비움)"),
    ("최종밴드단계", 10, "실험집단: 8주차 사용 색상"),
    ("나이", 6, "만 나이(세)"),
    ("신장", 7, "cm"),
    ("체중", 7, "kg"),
    ("KFES_사전", 8, "총점 16–64"),
    ("KFES_사후", 8, "총점 16–64"),
    ("의자일어서기_사전", 9, "30초 횟수(회)"),
    ("의자일어서기_사후", 9, "30초 횟수(회)"),
    ("팔굽히기_사전", 9, "30초 횟수(회)"),
    ("팔굽히기_사후", 9, "30초 횟수(회)"),
    ("2분제자리걷기_사전", 9, "2분 횟수(회)"),
    ("2분제자리걷기_사후", 9, "2분 횟수(회)"),
    ("등뒤로손잡기_사전", 9, "cm, 2회 중 최고값(겹치면 +)"),
    ("등뒤로손잡기_사후", 9, "cm, 2회 중 최고값(겹치면 +)"),
    ("앉아윗몸앞으로굽히기_사전", 10, "cm, 2회 중 최고값(지나면 +)"),
    ("앉아윗몸앞으로굽히기_사후", 10, "cm, 2회 중 최고값(지나면 +)"),
    ("2.44m갔다오기_사전", 9, "초, 소수 첫째 자리, 2회 중 최고값"),
    ("2.44m갔다오기_사후", 9, "초, 소수 첫째 자리, 2회 중 최고값"),
    ("측정자_사전", 8, "측정자 이름 또는 기호"),
    ("측정자_사후", 8, "측정자 이름 또는 기호"),
    ("비고", 18, "기록지 특이사항"),
]
N = 30
HDR_ROW, UNIT_ROW, FIRST = 3, 4, 5
LAST = FIRST + N - 1

wb = Workbook()
ws = wb.active
ws.title = "원자료"
ws["A1"] = "루프밴드 연구 원자료 (최초 참여자 30명 전원, 측정 기록지 원본 그대로 입력)"
ws["A1"].font = F_TITLE
ws["A2"] = "노란 칸에 입력. 열 이름(3행)은 분석 프로그램이 읽으므로 바꾸지 마세요. 빈칸은 결측으로 처리됩니다."
ws["A2"].font = F_NOTE

for c, (name, width, unit) in enumerate(COLS, 1):
    h = ws.cell(HDR_ROW, c, name)
    h.font, h.fill, h.alignment, h.border = F_HDR, FILL_HDR, CENTER, BOX
    u = ws.cell(UNIT_ROW, c, unit)
    u.font, u.alignment, u.border = F_NOTE, CENTER, BOX
    if name.endswith("_사전"):
        u.fill = FILL_PRE
    elif name.endswith("_사후"):
        u.fill = FILL_POST
    ws.column_dimensions[L(c)].width = width
    for r in range(FIRST, LAST + 1):
        cell = ws.cell(r, c)
        cell.font, cell.fill, cell.border = F_BASE, FILL_IN, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[HDR_ROW].height = 42
ws.row_dimensions[UNIT_ROW].height = 48
ws.freeze_panes = ws.cell(FIRST, 2)

col = {name: L(i) for i, (name, _, _) in enumerate(COLS, 1)}
for name, items in (("집단", "실험,통제"), ("이용층", "2층,4층"), ("최종분석포함", "예,아니오"),
                    ("제외시점", "사전검사 후,중재 중,사후검사 전")):
    dv = DataValidation(type="list", formula1=f'"{items}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col[name]}{FIRST}:{col[name]}{LAST}")
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="24", allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{col['출석횟수']}{FIRST}:{col['출석횟수']}{LAST}")

g = wb.create_sheet("안내")
guide = [
    ("이 파일의 목적", True),
    ("심사자 B가 요구한 '최초 참여자 30명 전원의 사전·사후 개인별 원자료'를 측정 기록지에서 그대로 옮겨 적는 양식입니다. "
     "1차·2차 재심본의 표 수치를 거꾸로 맞추는 용도가 아닙니다. 기록지에 적힌 값만 입력하세요.", False),
    ("입력 원칙", True),
    ("1. 두 사람이 따로 입력한 뒤 두 파일을 비교해 다른 칸을 기록지로 다시 확인합니다(이중 입력).", False),
    ("2. 기록지에 없는 값은 비워 둡니다. 추정하거나 다른 시점 값을 복사하지 않습니다.", False),
    ("3. SFT 기록 단위는 Rikli & Jones(2013) 기준을 따릅니다. 2.44m 갔다 오기는 0.1초 단위, 유연성 두 항목은 2회 중 최고값입니다. "
     "실제 기록지가 다른 단위(예: 정수 초)라면 그대로 입력하고 원고 측정방법에 그 사실을 적어야 합니다.", False),
    ("4. 제외자도 사전값이 있으면 입력합니다. 최종분석포함=아니오로 두면 분석에서 빠지고 흐름도 인원에만 쓰입니다.", False),
    ("5. 출석 기준(22회 이상)을 연구 시작 전에 정했는지는 원고에 따로 밝혀야 합니다.", False),
    ("분석", True),
    ("python3 loopband/tools/analyze.py <이 파일.xlsx> <결과.xlsx>  → 표 1, 사전 동질성, 낙상효능감·SFT 표, 공분산분석, 정규성, 자료 점검 시트가 만들어집니다.", False),
    ("SPSS로 분석할 경우에도 이 파일을 그대로 불러 쓰고, 출력물(.spv)을 함께 보관하세요. 심사자가 통계 출력물을 요구했습니다.", False),
]
g.column_dimensions["A"].width = 110
for i, (text, bold) in enumerate(guide, 1):
    c = g.cell(i, 1, text)
    c.font = F_BOLD if bold else F_BASE
    c.alignment = WRAP
wb.move_sheet("안내", offset=-1)
wb.save(OUT)
print("saved", OUT)
